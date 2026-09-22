#!/usr/bin/env python3
"""Plot the outer-radius history supplied in Attachment 2."""

from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from openpyxl import load_workbook


def load_radius_history(path: Path) -> tuple[np.ndarray, np.ndarray]:
    """Return strictly ordered time (s) and outer radius (cm)."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        headers = tuple(cell.value for cell in next(worksheet.iter_rows(max_row=1)))
        if headers[:2] != ("时间", "半径"):
            raise ValueError(f"unexpected Attachment 2 headers: {headers[:2]}")
        rows = [row[:2] for row in worksheet.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()

    if not rows or any(time is None or radius is None for time, radius in rows):
        raise ValueError("Attachment 2 contains missing time or radius values")
    times_s = np.asarray([time for time, _ in rows], dtype=float)
    radii_cm = np.asarray([radius for _, radius in rows], dtype=float)
    if not np.all(np.isfinite(times_s)) or not np.all(np.isfinite(radii_cm)):
        raise ValueError("Attachment 2 contains non-finite values")
    if np.any(np.diff(times_s) <= 0.0):
        raise ValueError("Attachment 2 times must increase strictly")
    if np.any(radii_cm <= 0.0):
        raise ValueError("Attachment 2 radii must be positive")
    return times_s, radii_cm


def configure_chinese_font() -> None:
    candidates = (
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    )
    for path in candidates:
        if path.exists():
            font_manager.fontManager.addfont(str(path))
            plt.rcParams["font.family"] = font_manager.FontProperties(
                fname=str(path)
            ).get_name()
            break
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["mathtext.fontset"] = "stix"


def draw_radius_history(times_s: np.ndarray, radii_cm: np.ndarray, output_stem: Path) -> None:
    configure_chinese_font()
    times_h = times_s / 3600.0

    figure, axis = plt.subplots(figsize=(10.0, 5.8), constrained_layout=True)
    figure.patch.set_facecolor("white")
    axis.set_facecolor("white")

    line_color = "#2463A8"
    axis.plot(
        times_h,
        radii_cm,
        color=line_color,
        linewidth=2.2,
        zorder=2,
        label="外半径",
    )
    axis.scatter(
        times_h,
        radii_cm,
        s=15,
        facecolor="white",
        edgecolor=line_color,
        linewidth=0.8,
        zorder=3,
        label="附件原始测点",
    )

    axis.annotate(
        f"初始：{radii_cm[0]:.3f} cm",
        xy=(times_h[0], radii_cm[0]),
        xytext=(5.0, radii_cm[0] - 0.035),
        arrowprops={"arrowstyle": "->", "color": "#3F4752", "lw": 1.0},
        color="#202631",
        fontsize=10.5,
    )
    axis.annotate(
        f"末值：{radii_cm[-1]:.3f} cm",
        xy=(times_h[-1], radii_cm[-1]),
        xytext=(52.0, radii_cm[-1] + 0.075),
        arrowprops={"arrowstyle": "->", "color": "#3F4752", "lw": 1.0},
        color="#202631",
        fontsize=10.5,
    )

    axis.set_title("附件 2：药材外半径随时间变化", fontsize=16, pad=13, weight="bold")
    axis.set_xlabel("时间 $t$ / h", fontsize=12)
    axis.set_ylabel("外半径 $R(t)$ / cm", fontsize=12)
    axis.set_xlim(float(times_h[0]), float(times_h[-1]))
    axis.set_ylim(1.15, 2.05)
    axis.set_xticks(np.arange(0.0, 72.0 + 0.1, 12.0))
    axis.set_yticks(np.arange(1.2, 2.0 + 0.01, 0.1))
    axis.grid(True, which="major", color="#D7DCE2", linewidth=0.8, alpha=0.85)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["left"].set_color("#59616C")
    axis.spines["bottom"].set_color("#59616C")
    axis.tick_params(colors="#343B45", labelsize=10.5)
    axis.legend(frameon=False, loc="upper right", ncol=2, fontsize=10.5)

    figure.text(
        0.995,
        0.006,
        f"数据来源：附件2.xlsx；采样间隔 30 min；共 {times_s.size} 个测点",
        ha="right",
        va="bottom",
        fontsize=8.8,
        color="#5E6670",
    )

    output_stem.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_stem.with_suffix(".png"), dpi=300, bbox_inches="tight")
    figure.savefig(output_stem.with_suffix(".svg"), bbox_inches="tight")
    plt.close(figure)


def main() -> None:
    root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=root / "A题" / "附件" / "附件2.xlsx",
    )
    parser.add_argument(
        "--output-stem",
        type=Path,
        default=root / "A题" / "figures" / "附件2外半径随时间变化趋势",
    )
    arguments = parser.parse_args()
    times_s, radii_cm = load_radius_history(arguments.input)
    draw_radius_history(times_s, radii_cm, arguments.output_stem)


if __name__ == "__main__":
    main()
