from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Patch
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent / "problem3" / "result3.xlsx"
OUT_PNG = HERE / "q3_3d_moisture_threshold.png"
OUT_PDF = HERE / "q3_3d_moisture_threshold.pdf"
CDRY = 0.15
MICROSOFT_YAHEI = HERE / "fonts" / "Microsoft YaHei.ttf"


def research_moist_cmap():
    """数模论文用清爽连续色带：低含水率浅金，高含水率深蓝。"""
    colors = [
        "#F7E6A1",
        "#B9DFA8",
        "#69C6B8",
        "#3A9CCB",
        "#3567A5",
        "#253B73",
    ]
    return LinearSegmentedColormap.from_list("cumcm_moisture", colors, N=256)


def load_field(path: Path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    radius_cm = np.asarray(headers[1:], dtype=float)
    times_s = []
    values = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        times_s.append(float(row[0]))
        values.append([float(v) for v in row[1:]])
    return np.asarray(times_s) / 3600.0, radius_cm, np.asarray(values)


def threshold_crossing_times(t_h, moisture, threshold):
    crossings = np.full(moisture.shape[1], np.nan)
    for j in range(moisture.shape[1]):
        y = moisture[:, j]
        hit = np.flatnonzero(y <= threshold)
        if hit.size == 0:
            continue
        i = int(hit[0])
        if i == 0:
            crossings[j] = t_h[0]
            continue
        t0, t1 = t_h[i - 1], t_h[i]
        y0, y1 = y[i - 1], y[i]
        if y1 == y0:
            crossings[j] = t1
        else:
            crossings[j] = t0 + (threshold - y0) * (t1 - t0) / (y1 - y0)
    return crossings


t_h, r_cm, moisture = load_field(SOURCE)
cross_t = threshold_crossing_times(t_h, moisture, CDRY)

# Display the complete process from 0 h.  PCHIP is used only to refine the
# rendering mesh: it passes through the workbook values and preserves monotone
# segments, avoiding the stair-step threshold boundary caused by 21 radial
# samples.  No model value or drying criterion is changed.
t_plot = np.linspace(float(t_h[0]), float(t_h[-1]), 720)
moisture_time_refined = PchipInterpolator(t_h, moisture, axis=0)(t_plot)
r_plot = np.linspace(float(r_cm.min()), float(r_cm.max()), 161)
c_plot = PchipInterpolator(r_cm, moisture_time_refined, axis=1)(r_plot)
c_plot = np.maximum(c_plot, np.finfo(float).tiny)
cross_t_plot = threshold_crossing_times(t_plot, c_plot, CDRY)
T, R = np.meshgrid(t_plot, r_plot, indexing="ij")
log_c_plot = np.log10(c_plot)

if MICROSOFT_YAHEI.exists():
    font_manager.fontManager.addfont(str(MICROSOFT_YAHEI))
    text_font_family = font_manager.FontProperties(fname=str(MICROSOFT_YAHEI)).get_name()
else:
    raise FileNotFoundError(f"缺少微软雅黑字体文件：{MICROSOFT_YAHEI}")

plt.rcParams.update(
    {
        "font.family": text_font_family,
        "mathtext.fontset": "stix",
        "axes.unicode_minus": False,
        "font.size": 10.5,
        "axes.linewidth": 0.8,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    }
)

fig = plt.figure(figsize=(8.1, 5.8))
ax = fig.add_subplot(111, projection="3d", computed_zorder=False)

moisture_cmap = research_moist_cmap()
threshold_face = "#A8AFB6"
threshold_edge = "#5F6871"
intersection_color = "#8B2F36"

# Split the moisture surface at the drying threshold.  Drawing the sub-threshold
# part first, the translucent threshold plane second, and the super-threshold
# part last gives the plane a genuine "cut through the surface" appearance.
below_threshold = np.ma.masked_where(c_plot > CDRY, log_c_plot)
above_threshold = np.ma.masked_where(c_plot < CDRY, log_c_plot)

ax.plot_surface(
    T,
    R,
    below_threshold,
    cmap=moisture_cmap,
    vmin=np.log10(0.04),
    vmax=np.log10(2.55),
    rcount=min(720, len(t_plot)),
    ccount=len(r_plot),
    linewidth=0,
    antialiased=True,
    alpha=0.72,
    shade=True,
    zorder=1,
)

plane_t, plane_r = np.meshgrid(
    np.linspace(0.0, 58.0, 59),
    np.linspace(float(r_cm.min()), float(r_cm.max()), 81),
    indexing="ij",
)
ax.plot_surface(
    plane_t,
    plane_r,
    np.full_like(plane_t, np.log10(CDRY)),
    color=threshold_face,
    alpha=0.25,
    linewidth=0,
    shade=False,
    antialiased=False,
    zorder=2,
)

ax.plot_surface(
    T,
    R,
    above_threshold,
    cmap=moisture_cmap,
    vmin=np.log10(0.04),
    vmax=np.log10(2.55),
    rcount=min(720, len(t_plot)),
    ccount=len(r_plot),
    linewidth=0,
    antialiased=True,
    alpha=0.98,
    shade=True,
    zorder=4,
)

# Trace the actual intersection C(t, r) = C_dry on the refined rendering mesh.
valid_crossing = np.isfinite(cross_t_plot)
ax.plot(
    cross_t_plot[valid_crossing],
    r_plot[valid_crossing],
    np.full(np.count_nonzero(valid_crossing), np.log10(CDRY)),
    color=intersection_color,
    linewidth=1.35,
    linestyle="--",
    alpha=0.98,
    solid_capstyle="round",
    zorder=6,
)

ax.set_xlim(0.0, 58.0)
ax.set_ylim(0.0, 2.0)
ax.set_zlim(np.log10(0.04), np.log10(2.60))
ax.set_xticks([0, 12, 24, 36, 48, 57])
ax.set_yticks(np.arange(0.0, 2.01, 0.5))
concentration_ticks = np.asarray([0.05, 0.10, 0.15, 0.30, 0.60, 1.20, 2.55])
ax.set_zticks(np.log10(concentration_ticks))
ax.set_zticklabels(["0.05", "0.10", "0.15", "0.30", "0.60", "1.20", "2.55"])
ax.set_xlabel(r"时间 $t$ / h", labelpad=12)
ax.set_ylabel(r"$r$ / cm", labelpad=12)
ax.set_zlabel(r"水分浓度 $C$ / (kg/kg)", labelpad=11)

ax.view_init(elev=25, azim=-64)
ax.set_box_aspect((1.75, 1.0, 0.92))
ax.grid(True)
for axis in (ax.xaxis, ax.yaxis, ax.zaxis):
    axis.pane.set_facecolor((1.0, 1.0, 1.0, 0.0))
    axis.pane.set_edgecolor((0.62, 0.62, 0.62, 1.0))
    axis._axinfo["grid"]["color"] = (0.76, 0.76, 0.76, 0.55)
    axis._axinfo["grid"]["linewidth"] = 0.55

legend_handles = [
    Patch(facecolor=moisture_cmap(0.72), edgecolor="none", alpha=0.97, label=r"$C(t,r)$"),
    Patch(
        facecolor=threshold_face,
        edgecolor=threshold_edge,
        alpha=0.32,
        label=r"$C_{\mathrm{dry}}=0.15\,\mathrm{kg/kg}$",
    ),
]
ax.legend(
    handles=legend_handles,
    loc="upper left",
    bbox_to_anchor=(0.02, 0.98),
    frameon=False,
    fontsize=9.6,
    handlelength=2.4,
    labelspacing=0.65,
)

fig.subplots_adjust(left=0.02, right=0.96, bottom=0.03, top=0.98)
fig.savefig(OUT_PNG, dpi=600, bbox_inches="tight", facecolor="white")
fig.savefig(OUT_PDF, bbox_inches="tight", facecolor="white")
plt.close(fig)

print("threshold crossings (h):")
for radius, crossing in zip(r_cm, cross_t):
    print(f"r={radius:.1f} cm: {crossing:.5f}")
