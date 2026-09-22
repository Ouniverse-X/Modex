#!/usr/bin/env python3
"""Run and archive the true-density/non-affine Q4 comparison experiment.

The script never overwrites an existing experiment.  A partially completed
suite can be continued explicitly with ``--resume``.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import platform
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from true_density_nonaffine import (
    NumericalConfig,
    SimulationResult,
    load_environment,
    load_radius,
    solve_model,
)


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DEFAULT_OUTPUT = HERE / "experiments" / "true_density_nonaffine_20260911_phase1"
ENVIRONMENT_PATH = ROOT / "A题" / "附件" / "附件1.xlsx"
RADIUS_PATH = ROOT / "A题" / "附件" / "附件2.xlsx"

BASE_CASES: dict[str, NumericalConfig] = {
    "space_n380_dt1": NumericalConfig(
        radial_cells=380, time_step_s=1.0, output_interval_s=60.0,
        maximum_time_h=60.0,
    ),
    "space_n760_dt1": NumericalConfig(
        radial_cells=760, time_step_s=1.0, output_interval_s=60.0,
        maximum_time_h=60.0,
    ),
    "reference_n1520_dt1": NumericalConfig(
        radial_cells=1520, time_step_s=1.0, output_interval_s=60.0,
        maximum_time_h=60.0,
    ),
    "time_n1520_dt4": NumericalConfig(
        radial_cells=1520, time_step_s=4.0, output_interval_s=60.0,
        maximum_time_h=60.0,
    ),
    "time_n1520_dt2": NumericalConfig(
        radial_cells=1520, time_step_s=2.0, output_interval_s=60.0,
        maximum_time_h=60.0,
    ),
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_default(value: Any) -> Any:
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, Path):
        return str(value)
    raise TypeError(f"cannot serialize {type(value)!r}")


def save_result(result: SimulationResult, case_dir: Path, case_name: str) -> dict[str, Any]:
    case_dir.mkdir(parents=True, exist_ok=False)
    np.savez_compressed(
        case_dir / "result_full_precision.npz",
        times_s=result.times_s,
        dry_mass_coordinates=result.dry_mass_coordinates,
        temperatures_c=result.temperatures_c,
        moistures_kg_kg=result.moistures_kg_kg,
        material_radii_m=result.material_radii_m,
        measured_radii_m=result.measured_radii_m,
        inferred_lengths_m=result.inferred_lengths_m,
        surface_temperatures_c=result.surface_temperatures_c,
        surface_moistures_kg_kg=result.surface_moistures_kg_kg,
        maximum_moistures_kg_kg=result.maximum_moistures_kg_kg,
    )
    metadata: dict[str, Any] = {
        "case": case_name,
        "config": asdict(result.config),
        "drying_crossing_s": result.drying_crossing_s,
        "drying_crossing_h": result.drying_crossing_s / 3600.0,
        "runtime_s": result.runtime_s,
        "diagnostics": result.diagnostics,
        "output_rows": int(result.times_s.size),
        "inferred_length_min_m": float(np.min(result.inferred_lengths_m)),
        "inferred_length_max_m": float(np.max(result.inferred_lengths_m)),
        "inferred_length_at_crossing_m": float(result.inferred_lengths_m[-1]),
    }
    (case_dir / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, default=_json_default) + "\n",
        encoding="utf-8",
    )
    return metadata


def load_archived_result(case_dir: Path) -> dict[str, Any]:
    archive = np.load(case_dir / "result_full_precision.npz")
    result = {name: archive[name] for name in archive.files}
    result["metadata"] = json.loads((case_dir / "metadata.json").read_text(encoding="utf-8"))
    return result


def run_case(case_name: str, config_dict: dict[str, Any], output_dir: str) -> dict[str, Any]:
    case_dir = Path(output_dir) / case_name
    config = NumericalConfig(**config_dict)
    environment = load_environment(ENVIRONMENT_PATH)
    radius = load_radius(RADIUS_PATH)
    result = solve_model(environment, radius, config, progress=False)
    return save_result(result, case_dir, case_name)


def interpolate_profiles(result: dict[str, Any], sample_times: np.ndarray, field: str) -> np.ndarray:
    source_times = result["times_s"]
    source = result[field]
    output = np.empty((sample_times.size, source.shape[1]))
    for column in range(source.shape[1]):
        output[:, column] = np.interp(sample_times, source_times, source[:, column])
    return output


def compare_cases(coarse: dict[str, Any], fine: dict[str, Any]) -> dict[str, float]:
    end_time = min(
        float(coarse["metadata"]["drying_crossing_s"]),
        float(fine["metadata"]["drying_crossing_s"]),
    )
    common_times = np.arange(0.0, math.floor(end_time / 60.0) * 60.0 + 0.1, 60.0)
    window_start = max(0.0, end_time - 6.0 * 3600.0)
    window = common_times >= window_start
    coarse_c = interpolate_profiles(coarse, common_times, "moistures_kg_kg")
    fine_c = interpolate_profiles(fine, common_times, "moistures_kg_kg")
    coarse_t = interpolate_profiles(coarse, common_times, "temperatures_c")
    fine_t = interpolate_profiles(fine, common_times, "temperatures_c")
    coarse_l = np.interp(common_times, coarse["times_s"], coarse["inferred_lengths_m"])
    fine_l = np.interp(common_times, fine["times_s"], fine["inferred_lengths_m"])
    difference_c = np.abs(coarse_c - fine_c)
    return {
        "common_end_s": float(common_times[-1]),
        "moisture_max_abs_all": float(np.max(difference_c)),
        "moisture_max_abs_last_6h": float(np.max(difference_c[window])),
        "temperature_max_abs_c": float(np.max(np.abs(coarse_t - fine_t))),
        "length_max_abs_m": float(np.max(np.abs(coarse_l - fine_l))),
        "crossing_difference_s": abs(
            float(coarse["metadata"]["drying_crossing_s"])
            - float(fine["metadata"]["drying_crossing_s"])
        ),
    }


def observed_order(error_coarse_medium: float, error_medium_fine: float) -> float:
    if error_coarse_medium <= 0.0 or error_medium_fine <= 0.0:
        return float("nan")
    return math.log(error_coarse_medium / error_medium_fine, 2.0)


def style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    for column_cells in worksheet.columns:
        width = min(28, max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def make_workbook(
    output_path: Path,
    reference: dict[str, Any],
    certification: dict[str, Any],
    convergence: dict[str, Any],
    suite_metadata: dict[str, Any],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "结果摘要"
    summary.append(["指标", "数值", "单位/说明"])
    ref_meta = reference["metadata"]
    cert_meta = certification["metadata"]
    rows = [
        ("模型", "真实湿密度—实测半径—非仿射收缩", "干质量坐标"),
        ("参考网格", ref_meta["config"]["radial_cells"], "累计干质量区间"),
        ("参考时间步", ref_meta["config"]["time_step_s"], "s"),
        ("Cmax=0.15 连续临界时刻", ref_meta["drying_crossing_s"], "s"),
        ("Cmax=0.15 连续临界时刻", ref_meta["drying_crossing_h"], "h"),
        ("四位小数安全阈值", suite_metadata["safety"]["numerical_target"], "kg/kg"),
        ("安全阈值连续到达时刻", cert_meta["drying_crossing_s"], "s"),
        ("建议整分钟安全时刻", suite_metadata["safety"]["safe_whole_minute_s"], "s"),
        ("建议整分钟安全时长", suite_metadata["safety"]["safe_whole_minute_h"], "h"),
        ("反演长度最小值", ref_meta["inferred_length_min_m"] * 100.0, "cm"),
        ("反演长度最大值", ref_meta["inferred_length_max_m"] * 100.0, "cm"),
        ("临界时刻反演长度", ref_meta["inferred_length_at_crossing_m"] * 100.0, "cm"),
        ("局部干质量最大相对误差", ref_meta["diagnostics"]["max_local_dry_mass_relative_error"], "-"),
        ("水分收支最大绝对残差", ref_meta["diagnostics"]["max_water_balance_abs"], "s^-1"),
    ]
    for row in rows:
        summary.append(row)
    style_sheet(summary)

    history = workbook.create_sheet("参考解时序")
    history.append([
        "time_s", "time_h", "measured_radius_cm", "inferred_length_cm",
        "length_ratio", "maximum_moisture", "center_moisture",
        "surface_moisture", "center_temperature_C", "surface_temperature_C",
        "max_nonaffine_deviation_mm",
    ])
    mu = reference["dry_mass_coordinates"]
    for i, time_s in enumerate(reference["times_s"]):
        affine = reference["measured_radii_m"][i] * np.sqrt(mu)
        deviation_mm = 1000.0 * np.max(np.abs(reference["material_radii_m"][i] - affine))
        history.append([
            float(time_s), float(time_s / 3600.0),
            float(100.0 * reference["measured_radii_m"][i]),
            float(100.0 * reference["inferred_lengths_m"][i]),
            float(reference["inferred_lengths_m"][i] / 0.25),
            float(reference["maximum_moistures_kg_kg"][i]),
            float(reference["moistures_kg_kg"][i, 0]),
            float(reference["surface_moistures_kg_kg"][i]),
            float(reference["temperatures_c"][i, 0]),
            float(reference["surface_temperatures_c"][i]),
            float(deviation_mm),
        ])
    style_sheet(history)

    profiles = workbook.create_sheet("参考解剖面")
    profiles.append([
        "time_s", "time_h", "dry_mass_coordinate", "material_radius_cm",
        "affine_radius_cm", "radial_deviation_mm", "temperature_C", "moisture_kg_kg",
    ])
    for i, time_s in enumerate(reference["times_s"]):
        affine = reference["measured_radii_m"][i] * np.sqrt(mu)
        for j, coordinate in enumerate(mu):
            profiles.append([
                float(time_s), float(time_s / 3600.0), float(coordinate),
                float(100.0 * reference["material_radii_m"][i, j]),
                float(100.0 * affine[j]),
                float(1000.0 * (reference["material_radii_m"][i, j] - affine[j])),
                float(reference["temperatures_c"][i, j]),
                float(reference["moistures_kg_kg"][i, j]),
            ])
    style_sheet(profiles)

    convergence_sheet = workbook.create_sheet("收敛性")
    convergence_sheet.append(["比较", "指标", "数值"])
    for comparison_name, values in convergence["comparisons"].items():
        for key, value in values.items():
            convergence_sheet.append([comparison_name, key, value])
    for key, value in convergence["orders"].items():
        convergence_sheet.append(["观测阶", key, value])
    style_sheet(convergence_sheet)

    environment = load_environment(ENVIRONMENT_PATH)
    environment_sheet = workbook.create_sheet("附件1输入")
    environment_sheet.append(["time_s", "temperature_C", "ambient_moisture"])
    for row in environment:
        environment_sheet.append([float(value) for value in row])
    style_sheet(environment_sheet)

    radius = load_radius(RADIUS_PATH)
    radius_sheet = workbook.create_sheet("附件2输入")
    radius_sheet.append(["time_s", "radius_cm"])
    for time_s, radius_m in radius:
        radius_sheet.append([float(time_s), float(100.0 * radius_m)])
    style_sheet(radius_sheet)
    workbook.save(output_path)


def make_figures(output_dir: Path, reference: dict[str, Any]) -> None:
    figure_dir = output_dir / "figures"
    figure_dir.mkdir(exist_ok=False)
    times_h = reference["times_s"] / 3600.0
    lengths_cm = 100.0 * reference["inferred_lengths_m"]
    mu = reference["dry_mass_coordinates"]
    nonaffine_mm = np.max(
        np.abs(reference["material_radii_m"] - reference["measured_radii_m"][:, None] * np.sqrt(mu)[None, :]),
        axis=1,
    ) * 1000.0

    chinese_font_path = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if chinese_font_path.exists():
        font_manager.fontManager.addfont(str(chinese_font_path))
        font_family = font_manager.FontProperties(fname=str(chinese_font_path)).get_name()
    else:
        font_family = "DejaVu Sans"
    plt.rcParams.update({
        "font.family": font_family,
        "axes.unicode_minus": False,
        "font.size": 15,
        "axes.labelsize": 17,
        "xtick.labelsize": 14,
        "ytick.labelsize": 14,
        "legend.fontsize": 15,
        "axes.grid": True,
        "grid.alpha": 0.25,
    })
    fig, axes = plt.subplots(2, 1, figsize=(8.0, 5.8), sharex=True, constrained_layout=True)
    axes[0].plot(times_h, lengths_cm, color="#d62728", linewidth=2.0)
    axes[0].axhline(25.0, color="0.4", linestyle="--", linewidth=1.0, label="初始长度")
    axes[0].set_ylabel("反演长度 / cm")
    axes[0].legend(frameon=False)
    axes[1].plot(times_h, nonaffine_mm, color="#2ca02c", linewidth=2.0)
    axes[1].set_ylabel("最大非仿射程度 / mm")
    axes[1].set_xlabel("时间 / h")
    fig.savefig(figure_dir / "geometry_diagnostics.png", dpi=300)
    fig.savefig(figure_dir / "geometry_diagnostics.svg")
    plt.close(fig)

    target_hours = [0.0, 6.0, 12.0, 24.0, 36.0, reference["metadata"]["drying_crossing_h"]]
    fig, axes = plt.subplots(1, 2, figsize=(10.0, 4.2), constrained_layout=True)
    colors = plt.cm.viridis(np.linspace(0.0, 1.0, len(target_hours)))
    for target_h, color in zip(target_hours, colors):
        index = int(np.argmin(np.abs(times_h - target_h)))
        radius_fraction = reference["material_radii_m"][index] / reference["measured_radii_m"][index]
        label = f"{times_h[index]:.2f} h"
        axes[0].plot(radius_fraction, reference["moistures_kg_kg"][index], color=color, label=label)
        axes[1].plot(radius_fraction, reference["temperatures_c"][index], color=color, label=label)
    axes[0].set_xlabel("Current normalized radius")
    axes[0].set_ylabel("Moisture (kg/kg, dry basis)")
    axes[1].set_xlabel("Current normalized radius")
    axes[1].set_ylabel("Temperature (deg C)")
    axes[1].legend(frameon=False, fontsize=8)
    fig.savefig(figure_dir / "field_profiles.png", dpi=300)
    fig.savefig(figure_dir / "field_profiles.svg")
    plt.close(fig)


def write_validation_report(
    output_dir: Path,
    cases: dict[str, dict[str, Any]],
    certification: dict[str, Any],
    convergence: dict[str, Any],
    safety: dict[str, Any],
) -> None:
    ref = cases["reference_n1520_dt1"]["metadata"]
    lines = [
        "# 真实湿密度—非仿射收缩对照实验验证报告",
        "",
        "## 主要结果",
        "",
        f"- 参考离散：$N=1520$，$\\Delta t=1\\,\\mathrm{{s}}$；",
        f"- $\\max C=0.15$ 的连续临界时刻：{ref['drying_crossing_s']:.9f} s = {ref['drying_crossing_h']:.12f} h；",
        f"- 用于四位小数安全判定的经验数值误差界：{safety['field_error_bound']:.12e} kg/kg；",
        f"- 安全判定数值阈值：{safety['numerical_target']:.12f} kg/kg；",
        f"- 建议整分钟安全时刻：{safety['safe_whole_minute_s']:.0f} s = {safety['safe_whole_minute_h']:.12f} h；",
        f"- 反演长度范围：{100.0 * ref['inferred_length_min_m']:.6f}–{100.0 * ref['inferred_length_max_m']:.6f} cm；",
        f"- 临界时刻反演长度：{100.0 * ref['inferred_length_at_crossing_m']:.6f} cm。",
        "",
        "## 守恒检查",
        "",
        f"- 每个材料环层干质量最大相对误差：{ref['diagnostics']['max_local_dry_mass_relative_error']:.3e}；",
        f"- 水分方程最大绝对收支残差：{ref['diagnostics']['max_water_balance_abs']:.3e} s^-1；",
        f"- 水分方程最大相对收支残差：{ref['diagnostics']['max_water_balance_rel']:.3e}；",
        f"- 全程最大含水率位于中心：{ref['diagnostics']['maximum_moisture_at_center']}；",
        f"- 径向含水率单调性最大违背量：{ref['diagnostics']['max_radial_monotonicity_violation']:.3e}。",
        "",
        "## 网格和时间步收敛",
        "",
        "| 比较 | 临界时刻差/s | 末6 h含水率最大差 | 全程含水率最大差 | 温度最大差/°C | 长度最大差/m |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for name, item in convergence["comparisons"].items():
        lines.append(
            f"| {name} | {item['crossing_difference_s']:.9f} | "
            f"{item['moisture_max_abs_last_6h']:.9e} | {item['moisture_max_abs_all']:.9e} | "
            f"{item['temperature_max_abs_c']:.9e} | {item['length_max_abs_m']:.9e} |"
        )
    lines.extend([
        "",
        f"空间观测阶（末6 h含水率误差）：{convergence['orders']['space_moisture_last_6h']:.6f}；",
        f"时间观测阶（末6 h含水率误差）：{convergence['orders']['time_moisture_last_6h']:.6f}。",
        "",
        "## 物理解释提醒",
        "",
        "附件2半径被作为强制外边界，附录4密度被作为真实局部湿体积密度，干质量逐环严格守恒。"
        "因此长度是由相容条件反演的输出，不是人为给定。若反演长度显著增大，并非数值程序自动证明药材真实伸长，"
        "而是说明这两组数据与当前‘无额外孔隙/结构变量’闭合关系共同要求相应的轴向补偿变形。",
        "",
        "安全误差界由最细两层空间网格差与最细两层时间步差之和再乘 1.5 得到，是保守的数值收敛估计，"
        "不是严格的数学后验误差定理。",
    ])
    (output_dir / "validation_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def summarize(output_dir: Path) -> dict[str, Any]:
    cases = {name: load_archived_result(output_dir / name) for name in BASE_CASES}
    comparisons = {
        "space_380_vs_760": compare_cases(cases["space_n380_dt1"], cases["space_n760_dt1"]),
        "space_760_vs_1520": compare_cases(cases["space_n760_dt1"], cases["reference_n1520_dt1"]),
        "time_4s_vs_2s": compare_cases(cases["time_n1520_dt4"], cases["time_n1520_dt2"]),
        "time_2s_vs_1s": compare_cases(cases["time_n1520_dt2"], cases["reference_n1520_dt1"]),
    }
    orders = {
        "space_moisture_last_6h": observed_order(
            comparisons["space_380_vs_760"]["moisture_max_abs_last_6h"],
            comparisons["space_760_vs_1520"]["moisture_max_abs_last_6h"],
        ),
        "time_moisture_last_6h": observed_order(
            comparisons["time_4s_vs_2s"]["moisture_max_abs_last_6h"],
            comparisons["time_2s_vs_1s"]["moisture_max_abs_last_6h"],
        ),
        "space_crossing": observed_order(
            comparisons["space_380_vs_760"]["crossing_difference_s"],
            comparisons["space_760_vs_1520"]["crossing_difference_s"],
        ),
        "time_crossing": observed_order(
            comparisons["time_4s_vs_2s"]["crossing_difference_s"],
            comparisons["time_2s_vs_1s"]["crossing_difference_s"],
        ),
    }
    convergence = {"comparisons": comparisons, "orders": orders}

    field_error_bound = 1.5 * (
        comparisons["space_760_vs_1520"]["moisture_max_abs_last_6h"]
        + comparisons["time_2s_vs_1s"]["moisture_max_abs_last_6h"]
    )
    numerical_target = 0.14995 - field_error_bound
    if numerical_target <= 0.149:
        raise RuntimeError(f"convergence error bound is unexpectedly large: {field_error_bound}")
    certification_config = NumericalConfig(
        radial_cells=1520,
        time_step_s=1.0,
        output_interval_s=60.0,
        maximum_time_h=60.0,
        drying_threshold_kg_kg=numerical_target,
    )
    certification_dir = output_dir / "certification_n1520_dt1"
    if certification_dir.exists():
        certification = load_archived_result(certification_dir)
    else:
        result = solve_model(
            load_environment(ENVIRONMENT_PATH), load_radius(RADIUS_PATH), certification_config
        )
        save_result(result, certification_dir, "certification_n1520_dt1")
        certification = load_archived_result(certification_dir)
    certified_crossing_s = float(certification["metadata"]["drying_crossing_s"])
    safe_whole_minute_s = math.ceil(certified_crossing_s / 60.0) * 60.0
    safety = {
        "field_error_bound": field_error_bound,
        "display_limit": 0.14995,
        "numerical_target": numerical_target,
        "certified_target_crossing_s": certified_crossing_s,
        "safe_whole_minute_s": safe_whole_minute_s,
        "safe_whole_minute_h": safe_whole_minute_s / 3600.0,
        "qualification": "conservative convergence estimate, not a rigorous a posteriori theorem",
    }
    suite_metadata: dict[str, Any] = {
        "model": "true wet density + measured outer radius + non-affine radial mapping",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "python": sys.version,
        "platform": platform.platform(),
        "cpu_count": os.cpu_count(),
        "inputs": {
            "environment": {"path": str(ENVIRONMENT_PATH), "sha256": sha256_file(ENVIRONMENT_PATH)},
            "radius": {"path": str(RADIUS_PATH), "sha256": sha256_file(RADIUS_PATH)},
            "solver": {"path": str(HERE / 'true_density_nonaffine.py'), "sha256": sha256_file(HERE / 'true_density_nonaffine.py')},
            "runner": {"path": str(Path(__file__).resolve()), "sha256": sha256_file(Path(__file__).resolve())},
        },
        "cases": {name: value["metadata"] for name, value in cases.items()},
        "certification_case": certification["metadata"],
        "convergence": convergence,
        "safety": safety,
    }
    (output_dir / "run_metadata.json").write_text(
        json.dumps(suite_metadata, ensure_ascii=False, indent=2, default=_json_default) + "\n",
        encoding="utf-8",
    )
    make_workbook(
        output_dir / "result4-2.xlsx",
        cases["reference_n1520_dt1"],
        certification,
        convergence,
        suite_metadata,
    )
    make_figures(output_dir, cases["reference_n1520_dt1"])
    write_validation_report(output_dir, cases, certification, convergence, safety)
    return suite_metadata


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--workers", type=int, default=5)
    parser.add_argument("--resume", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_dir = args.output.resolve()
    if output_dir.exists() and not args.resume:
        raise FileExistsError(
            f"experiment directory already exists: {output_dir}; use --resume only for this suite"
        )
    output_dir.mkdir(parents=True, exist_ok=args.resume)
    pending = {
        name: config for name, config in BASE_CASES.items()
        if not (output_dir / name / "result_full_precision.npz").exists()
    }
    if pending:
        with ProcessPoolExecutor(max_workers=min(args.workers, len(pending))) as pool:
            futures = {
                pool.submit(run_case, name, asdict(config), str(output_dir)): name
                for name, config in pending.items()
            }
            for future in as_completed(futures):
                name = futures[future]
                metadata = future.result()
                print(
                    f"[{name}] crossing={metadata['drying_crossing_h']:.12f} h; "
                    f"runtime={metadata['runtime_s']:.2f} s",
                    flush=True,
                )
    metadata = summarize(output_dir)
    print(json.dumps(metadata["safety"], ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
