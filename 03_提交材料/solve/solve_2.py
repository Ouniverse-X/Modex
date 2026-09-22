#!/usr/bin/env python3
"""Solve Problem A, Question 2 with a coupled conservative radial model.

Model
-----
    rho(C) cp(C) dT/dt = (1/r) d/dr (r k(C) dT/dr)
    dC/dt              = (1/r) d/dr (r D(C,T_K) dC/dr)

The cylinder radius is fixed at 0.02 m in Question 2.  The environment data are
piecewise-linearly interpolated from Attachment 1.  Space is discretized by a
node-centred radial finite-volume method with surface refinement; the coupled
method-of-lines system is integrated by SciPy's variable-order implicit BDF
solver with sparse Jacobian structure and strict error tolerances.

The program writes the requested result2.xlsx, paper tables, full-precision CSV
files, a convergence/validation report, and reproducibility metadata.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import platform
import sys
import time
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Sequence

import numpy as np
import scipy
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from scipy.integrate import solve_ivp
from scipy.sparse import csr_matrix, lil_matrix


@dataclass(frozen=True)
class PhysicalParameters:
    radius_m: float = 0.02
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8.0e-7
    initial_temperature_c: float = 28.0
    initial_moisture_kg_kg: float = 2.55
    density_intercept_kg_m3: float = 650.0
    density_moisture_slope_kg_m3: float = 128.0
    heat_capacity_intercept_j_kg_k: float = 1450.0
    heat_capacity_moisture_scale_j_kg_k: float = 2736.0
    conductivity_intercept_w_m_k: float = 0.21
    conductivity_moisture_scale_w_m_k: float = 0.38
    diffusivity_prefactor_m2_s: float = 2.4e-3
    diffusivity_moisture_exponent: float = 0.45
    diffusivity_temperature_exponent_k: float = 3850.0

    def evaluate_properties(
        self, moisture: np.ndarray, temperature_c: np.ndarray
    ) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
        moisture = np.asarray(moisture, dtype=float)
        temperature_c = np.asarray(temperature_c, dtype=float)
        if moisture.shape != temperature_c.shape:
            raise ValueError("temperature and moisture arrays must have equal shapes")
        if not np.all(np.isfinite(moisture)) or np.any(moisture <= 0.0):
            raise ValueError("moisture must be positive and finite")
        temperature_k = temperature_c + 273.15
        if not np.all(np.isfinite(temperature_k)) or np.any(temperature_k <= 0.0):
            raise ValueError("absolute temperature must be positive and finite")

        ratio = moisture / (moisture + 1.0)
        density = (
            self.density_intercept_kg_m3
            + self.density_moisture_slope_kg_m3 * moisture
        )
        heat_capacity = (
            self.heat_capacity_intercept_j_kg_k
            + self.heat_capacity_moisture_scale_j_kg_k * ratio
        )
        conductivity = (
            self.conductivity_intercept_w_m_k
            + self.conductivity_moisture_scale_w_m_k * ratio
        )
        diffusivity = self.diffusivity_prefactor_m2_s * np.exp(
            -self.diffusivity_moisture_exponent / moisture
            -self.diffusivity_temperature_exponent_k / temperature_k
        )
        if (
            np.any(density <= 0.0)
            or np.any(heat_capacity <= 0.0)
            or np.any(conductivity <= 0.0)
            or np.any(diffusivity <= 0.0)
            or not np.all(np.isfinite(diffusivity))
        ):
            raise ValueError("an empirical material property became non-positive")
        return density, heat_capacity, conductivity, diffusivity


@dataclass(frozen=True)
class NumericalConfig:
    radial_cells: int = 760
    end_time_s: float = 10800.0
    output_interval_s: float = 1.0
    surface_layer_m: float = 0.002
    surface_refinement_factor: int = 10
    relative_tolerance: float = 2.0e-9
    temperature_absolute_tolerance: float = 2.0e-9
    moisture_absolute_tolerance: float = 2.0e-11
    maximum_step_s: float = 5.0
    first_step_s: float = 1.0e-4

    def validate(self) -> None:
        if self.radial_cells < 38 or self.radial_cells % 19 != 0:
            raise ValueError("radial_cells must be a multiple of 19 and at least 38")
        if self.end_time_s <= 0.0 or self.output_interval_s <= 0.0:
            raise ValueError("time intervals must be positive")
        if abs(self.end_time_s / self.output_interval_s - round(
            self.end_time_s / self.output_interval_s
        )) > 1.0e-12:
            raise ValueError("end_time_s must be divisible by output_interval_s")
        if not 0.0 < self.surface_layer_m < 0.02:
            raise ValueError("surface_layer_m must lie inside the cylinder")
        if self.surface_refinement_factor < 1:
            raise ValueError("surface_refinement_factor must be at least one")
        for name, value in (
            ("relative_tolerance", self.relative_tolerance),
            ("temperature_absolute_tolerance", self.temperature_absolute_tolerance),
            ("moisture_absolute_tolerance", self.moisture_absolute_tolerance),
            ("maximum_step_s", self.maximum_step_s),
            ("first_step_s", self.first_step_s),
        ):
            if not math.isfinite(value) or value <= 0.0:
                raise ValueError(f"{name} must be positive and finite")


@dataclass(frozen=True)
class EnvironmentData:
    times_s: np.ndarray
    temperatures_c: np.ndarray
    moistures_kg_kg: np.ndarray

    def validate(self, end_time_s: float) -> None:
        if self.times_s.ndim != 1 or self.times_s.size < 2:
            raise ValueError("environment must contain at least two time points")
        if not (
            self.times_s.shape
            == self.temperatures_c.shape
            == self.moistures_kg_kg.shape
        ):
            raise ValueError("environment columns have inconsistent shapes")
        if not (
            np.all(np.isfinite(self.times_s))
            and np.all(np.isfinite(self.temperatures_c))
            and np.all(np.isfinite(self.moistures_kg_kg))
        ):
            raise ValueError("environment contains a non-finite value")
        if np.any(np.diff(self.times_s) <= 0.0):
            raise ValueError("environment times must be strictly increasing")
        if self.times_s[0] > 0.0 or self.times_s[-1] < end_time_s:
            raise ValueError(
                f"environment must cover [0, {end_time_s}], got "
                f"[{self.times_s[0]}, {self.times_s[-1]}]"
            )
        if np.any(self.moistures_kg_kg <= 0.0):
            raise ValueError("environment moisture must be positive")

    def interpolate(self, t_s: float) -> tuple[float, float]:
        temperature = float(np.interp(t_s, self.times_s, self.temperatures_c))
        moisture = float(np.interp(t_s, self.times_s, self.moistures_kg_kg))
        return temperature, moisture


@dataclass
class SolverDiagnostics:
    runtime_s: float
    success: bool
    message: str
    function_evaluations: int
    jacobian_evaluations: int
    lu_decompositions: int
    returned_time_points: int
    min_temperature_c: float
    max_temperature_c: float
    min_moisture_kg_kg: float
    max_moisture_kg_kg: float
    min_density_kg_m3: float
    max_density_kg_m3: float
    min_heat_capacity_j_kg_k: float
    max_heat_capacity_j_kg_k: float
    min_conductivity_w_m_k: float
    max_conductivity_w_m_k: float
    min_diffusivity_m2_s: float
    max_diffusivity_m2_s: float
    max_one_second_moisture_balance_abs: float
    max_one_second_moisture_balance_rel: float


@dataclass
class SimulationResult:
    times_s: np.ndarray
    radii_cm: np.ndarray
    temperatures_c: np.ndarray
    moistures_kg_kg: np.ndarray
    diagnostics: SolverDiagnostics
    config: NumericalConfig


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_environment(path: Path, end_time_s: float) -> EnvironmentData:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows or tuple(str(value).strip() for value in rows[0][:3]) != (
        "时间",
        "温度",
        "水分浓度",
    ):
        raise ValueError("附件1 must contain 时间、温度、水分浓度 columns")
    values: list[tuple[float, float, float]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) < 3 or any(value is None for value in row[:3]):
            raise ValueError(f"missing environment value in row {row_number}")
        try:
            values.append(tuple(float(value) for value in row[:3]))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"non-numeric environment value in row {row_number}") from exc
    array = np.asarray(values, dtype=float)
    environment = EnvironmentData(array[:, 0], array[:, 1], array[:, 2])
    environment.validate(end_time_s)
    return environment


def make_radial_grid(
    radius_m: float,
    radial_cells: int,
    surface_layer_m: float,
    surface_refinement_factor: int,
) -> np.ndarray:
    """Construct a two-zone grid refined in the outermost surface layer."""
    bulk_length = radius_m - surface_layer_m
    effective_length = bulk_length + surface_refinement_factor * surface_layer_m
    bulk_step = effective_length / radial_cells
    surface_step = bulk_step / surface_refinement_factor
    bulk_cells_float = bulk_length / bulk_step
    surface_cells_float = surface_layer_m / surface_step
    bulk_cells = int(round(bulk_cells_float))
    surface_cells = int(round(surface_cells_float))
    if (
        abs(bulk_cells_float - bulk_cells) > 1.0e-10
        or abs(surface_cells_float - surface_cells) > 1.0e-10
        or bulk_cells + surface_cells != radial_cells
    ):
        raise ValueError(
            "radial_cells is incompatible with the two-zone grid; "
            "use a multiple of 19 for the 2 mm / factor-10 refinement"
        )
    bulk = np.arange(bulk_cells + 1, dtype=float) * bulk_step
    surface = bulk_length + np.arange(1, surface_cells + 1, dtype=float) * surface_step
    radii = np.concatenate((bulk, surface))
    radii[0] = 0.0
    radii[-1] = radius_m
    if radii.size != radial_cells + 1 or np.any(np.diff(radii) <= 0.0):
        raise AssertionError("constructed radial grid is invalid")
    return radii


def radial_storage_weights(radii_m: np.ndarray) -> np.ndarray:
    faces = 0.5 * (radii_m[:-1] + radii_m[1:])
    weights = np.empty_like(radii_m)
    weights[0] = 0.5 * faces[0] ** 2
    weights[1:-1] = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    weights[-1] = 0.5 * (radii_m[-1] ** 2 - faces[-1] ** 2)
    return weights


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    denominator = left + right
    if np.any(left <= 0.0) or np.any(right <= 0.0) or np.any(denominator <= 0.0):
        raise ValueError("interface transport coefficients must be positive")
    return 2.0 * left * right / denominator


def make_jacobian_sparsity(node_count: int) -> csr_matrix:
    """Return the block-tridiagonal dependency pattern of the coupled system."""
    pattern = lil_matrix((2 * node_count, 2 * node_count), dtype=np.int8)
    for i in range(node_count):
        neighbours = range(max(0, i - 1), min(node_count, i + 2))
        for j in neighbours:
            pattern[i, j] = 1
            pattern[i, node_count + j] = 1
            pattern[node_count + i, j] = 1
            pattern[node_count + i, node_count + j] = 1
    return pattern.tocsr()


class CoupledRadialSystem:
    def __init__(
        self,
        environment: EnvironmentData,
        parameters: PhysicalParameters,
        radii_m: np.ndarray,
    ) -> None:
        self.environment = environment
        self.parameters = parameters
        self.radii_m = radii_m
        self.node_count = radii_m.size
        self.radius_m = float(radii_m[-1])
        self.faces_m = 0.5 * (radii_m[:-1] + radii_m[1:])
        self.spacings_m = np.diff(radii_m)
        self.weights_m2 = radial_storage_weights(radii_m)

    def transport_net(
        self,
        field: np.ndarray,
        gamma_nodes: np.ndarray,
        boundary_transfer: float,
        external_value: float,
    ) -> np.ndarray:
        gamma_faces = harmonic_mean(gamma_nodes[:-1], gamma_nodes[1:])
        interface_flux = (
            self.faces_m
            * gamma_faces
            * np.diff(field)
            / self.spacings_m
        )
        net = np.empty_like(field)
        net[0] = interface_flux[0]
        net[1:-1] = interface_flux[1:] - interface_flux[:-1]
        net[-1] = (
            self.radius_m * boundary_transfer * (external_value - field[-1])
            - interface_flux[-1]
        )
        return net

    def __call__(self, t_s: float, state: np.ndarray) -> np.ndarray:
        n = self.node_count
        temperature_c = state[:n]
        moisture = state[n:]
        density, heat_capacity, conductivity, diffusivity = (
            self.parameters.evaluate_properties(moisture, temperature_c)
        )
        external_temperature, external_moisture = self.environment.interpolate(t_s)
        heat_net = self.transport_net(
            temperature_c,
            conductivity,
            self.parameters.heat_transfer_w_m2_k,
            external_temperature,
        )
        moisture_net = self.transport_net(
            moisture,
            diffusivity,
            self.parameters.mass_transfer_m_s,
            external_moisture,
        )
        temperature_rate = heat_net / (
            self.weights_m2 * density * heat_capacity
        )
        moisture_rate = moisture_net / self.weights_m2
        return np.concatenate((temperature_rate, moisture_rate))


def requested_output_indices(radii_m: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    requested_cm = np.arange(21, dtype=float) * 0.1
    indices = np.empty(requested_cm.size, dtype=int)
    for j, radius_cm in enumerate(requested_cm):
        target_m = radius_cm / 100.0
        index = int(np.argmin(np.abs(radii_m - target_m)))
        if abs(float(radii_m[index]) - target_m) > 1.0e-12:
            raise ValueError(f"requested radius {radius_cm:g} cm is not a grid node")
        indices[j] = index
    return requested_cm, indices


def moisture_balance_diagnostic(
    environment: EnvironmentData,
    parameters: PhysicalParameters,
    weights: np.ndarray,
    times_s: np.ndarray,
    all_moistures: np.ndarray,
) -> tuple[float, float]:
    inventory = weights @ all_moistures
    surface = all_moistures[-1, :]
    external = np.interp(times_s, environment.times_s, environment.moistures_kg_kg)
    boundary_rate = (
        parameters.radius_m
        * parameters.mass_transfer_m_s
        * (external - surface)
    )
    dt = np.diff(times_s)
    trapezoidal_change = 0.5 * (boundary_rate[:-1] + boundary_rate[1:]) * dt
    residual = np.diff(inventory) - trapezoidal_change
    scale = np.maximum.reduce(
        (
            np.abs(np.diff(inventory)),
            np.abs(trapezoidal_change),
            np.full_like(residual, 1.0e-30),
        )
    )
    return float(np.max(np.abs(residual))), float(np.max(np.abs(residual) / scale))


def solve_model(
    environment: EnvironmentData,
    parameters: PhysicalParameters,
    config: NumericalConfig,
    progress_label: str | None = None,
    integration_method: str = "BDF",
) -> SimulationResult:
    if integration_method not in {"BDF", "Radau"}:
        raise ValueError("integration_method must be 'BDF' or 'Radau'")
    config.validate()
    environment.validate(config.end_time_s)
    radii_m = make_radial_grid(
        parameters.radius_m,
        config.radial_cells,
        config.surface_layer_m,
        config.surface_refinement_factor,
    )
    system = CoupledRadialSystem(environment, parameters, radii_m)
    node_count = radii_m.size
    state0 = np.concatenate(
        (
            np.full(node_count, parameters.initial_temperature_c, dtype=float),
            np.full(node_count, parameters.initial_moisture_kg_kg, dtype=float),
        )
    )
    output_times = np.arange(
        0.0,
        config.end_time_s + 0.5 * config.output_interval_s,
        config.output_interval_s,
        dtype=float,
    )
    absolute_tolerance = np.concatenate(
        (
            np.full(node_count, config.temperature_absolute_tolerance),
            np.full(node_count, config.moisture_absolute_tolerance),
        )
    )
    if progress_label:
        print(
            f"[{progress_label}] N={config.radial_cells}, "
            f"rtol={config.relative_tolerance:.1e}, "
            f"max_step={config.maximum_step_s:g} s",
            flush=True,
        )
    started = time.perf_counter()
    solution = solve_ivp(
        system,
        (0.0, config.end_time_s),
        state0,
        method=integration_method,
        t_eval=output_times,
        rtol=config.relative_tolerance,
        atol=absolute_tolerance,
        jac_sparsity=make_jacobian_sparsity(node_count),
        max_step=config.maximum_step_s,
        first_step=min(config.first_step_s, config.end_time_s),
    )
    runtime_s = time.perf_counter() - started
    if not solution.success:
        raise RuntimeError(
            f"{integration_method} integration failed: {solution.message}"
        )
    if solution.t.size != output_times.size or not np.array_equal(
        solution.t, output_times
    ):
        raise RuntimeError(
            f"{integration_method} solver did not return the requested output time grid"
        )
    if not np.all(np.isfinite(solution.y)):
        raise RuntimeError(f"{integration_method} solution contains non-finite values")

    all_temperature = solution.y[:node_count, :]
    all_moisture = solution.y[node_count:, :]
    if np.min(all_moisture) <= 0.0:
        raise RuntimeError("computed a non-positive moisture value")

    density, heat_capacity, conductivity, diffusivity = parameters.evaluate_properties(
        all_moisture.ravel(), all_temperature.ravel()
    )
    balance_abs, balance_rel = moisture_balance_diagnostic(
        environment,
        parameters,
        system.weights_m2,
        output_times,
        all_moisture,
    )
    radii_cm, indices = requested_output_indices(radii_m)
    result = SimulationResult(
        times_s=output_times[1:].astype(int),
        radii_cm=radii_cm,
        temperatures_c=all_temperature[indices, 1:].T.copy(),
        moistures_kg_kg=all_moisture[indices, 1:].T.copy(),
        diagnostics=SolverDiagnostics(
            runtime_s=runtime_s,
            success=bool(solution.success),
            message=str(solution.message),
            function_evaluations=int(solution.nfev),
            jacobian_evaluations=int(solution.njev),
            lu_decompositions=int(solution.nlu),
            returned_time_points=int(solution.t.size),
            min_temperature_c=float(np.min(all_temperature)),
            max_temperature_c=float(np.max(all_temperature)),
            min_moisture_kg_kg=float(np.min(all_moisture)),
            max_moisture_kg_kg=float(np.max(all_moisture)),
            min_density_kg_m3=float(np.min(density)),
            max_density_kg_m3=float(np.max(density)),
            min_heat_capacity_j_kg_k=float(np.min(heat_capacity)),
            max_heat_capacity_j_kg_k=float(np.max(heat_capacity)),
            min_conductivity_w_m_k=float(np.min(conductivity)),
            max_conductivity_w_m_k=float(np.max(conductivity)),
            min_diffusivity_m2_s=float(np.min(diffusivity)),
            max_diffusivity_m2_s=float(np.max(diffusivity)),
            max_one_second_moisture_balance_abs=balance_abs,
            max_one_second_moisture_balance_rel=balance_rel,
        ),
        config=config,
    )
    if progress_label:
        print(
            f"[{progress_label}] complete in {runtime_s:.2f} s; "
            f"nfev={solution.nfev}, nlu={solution.nlu}",
            flush=True,
        )
    return result


def compare_results(reference: SimulationResult, candidate: SimulationResult) -> dict:
    if not np.array_equal(reference.times_s, candidate.times_s) or not np.array_equal(
        reference.radii_cm, candidate.radii_cm
    ):
        raise ValueError("results have different output coordinates")

    def compare_field(reference_values: np.ndarray, candidate_values: np.ndarray) -> dict:
        difference = candidate_values - reference_values
        absolute = np.abs(difference)
        flat_index = int(np.argmax(absolute))
        time_index, radius_index = np.unravel_index(flat_index, absolute.shape)
        rounded_equal = np.round(reference_values, 4) == np.round(candidate_values, 4)
        return {
            "max_abs_difference": float(absolute[time_index, radius_index]),
            "rms_difference": float(np.sqrt(np.mean(difference * difference))),
            "max_location_time_s": int(reference.times_s[time_index]),
            "max_location_radius_cm": float(reference.radii_cm[radius_index]),
            "four_decimal_agreement_count": int(np.count_nonzero(rounded_equal)),
            "comparison_count": int(rounded_equal.size),
            "four_decimal_agreement_fraction": float(np.mean(rounded_equal)),
        }

    return {
        "temperature": compare_field(
            reference.temperatures_c, candidate.temperatures_c
        ),
        "moisture": compare_field(
            reference.moistures_kg_kg, candidate.moistures_kg_kg
        ),
    }


def add_richardson_estimates(
    coarse_to_medium: dict, medium_to_fine: dict
) -> dict:
    result: dict[str, dict] = {}
    for field in ("temperature", "moisture"):
        coarse_difference = coarse_to_medium[field]["max_abs_difference"]
        fine_difference = medium_to_fine[field]["max_abs_difference"]
        if coarse_difference > 0.0 and fine_difference > 0.0:
            order = math.log(coarse_difference / fine_difference, 2.0)
            denominator = 2.0**order - 1.0
            estimate = fine_difference / denominator if denominator > 0.0 else math.inf
        else:
            order = math.inf
            estimate = 0.0
        result[field] = {
            **medium_to_fine[field],
            "coarse_to_medium_max_abs_difference": coarse_difference,
            "observed_order_max_norm": order,
            "estimated_fine_grid_max_error": estimate,
            "estimated_fine_grid_max_error_with_1p25_safety": 1.25 * estimate,
        }
    return result


def write_full_precision_csv(
    path: Path,
    times_s: np.ndarray,
    radii_cm: np.ndarray,
    values: np.ndarray,
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["时间/s"] + [f"{radius:.1f} cm" for radius in radii_cm])
        for t_s, row in zip(times_s, values):
            writer.writerow([int(t_s)] + [f"{value:.12f}" for value in row])


def write_environment_csv(path: Path, environment: EnvironmentData, end_time_s: float) -> None:
    mask = environment.times_s <= end_time_s
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["时间/s", "温度/°C", "水分浓度/(kg/kg)"])
        for row in zip(
            environment.times_s[mask],
            environment.temperatures_c[mask],
            environment.moistures_kg_kg[mask],
        ):
            writer.writerow([f"{row[0]:.0f}", f"{row[1]:.15g}", f"{row[2]:.15g}"])


def write_result_xlsx(template_path: Path, output_path: Path, result: SimulationResult) -> None:
    workbook = load_workbook(template_path)
    try:
        if workbook.sheetnames != ["温度", "水分浓度"]:
            raise ValueError(f"unexpected template worksheets: {workbook.sheetnames}")
        for sheet_name, field in (
            ("温度", result.temperatures_c),
            ("水分浓度", result.moistures_kg_kg),
        ):
            worksheet = workbook[sheet_name]
            header_style = copy(worksheet["B1"]._style)
            time_style = copy(worksheet["A2"]._style)
            value_style = copy(worksheet["B2"]._style)
            if worksheet.max_row > 1:
                worksheet.delete_rows(2, worksheet.max_row - 1)
            worksheet["A1"] = "时间\\到药材中心的距离/cm"
            worksheet["A1"]._style = copy(header_style)
            worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
            for column_index, radius_cm in enumerate(result.radii_cm, start=2):
                cell = worksheet.cell(row=1, column=column_index, value=float(radius_cm))
                cell._style = copy(header_style)
                cell.number_format = "0.0"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row_index, (t_s, row) in enumerate(
                zip(result.times_s, field), start=2
            ):
                time_cell = worksheet.cell(row=row_index, column=1, value=int(t_s))
                time_cell._style = copy(time_style)
                for column_index, value in enumerate(row, start=2):
                    cell = worksheet.cell(
                        row=row_index,
                        column=column_index,
                        value=round(float(value), 4),
                    )
                    cell._style = copy(value_style)
                    cell.number_format = "0.0000"
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = f"A1:V{len(result.times_s) + 1}"
            worksheet.column_dimensions["A"].width = 24
            for column in range(2, 23):
                worksheet.column_dimensions[worksheet.cell(1, column).column_letter].width = 12
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = output_path.with_suffix(output_path.suffix + ".tmp")
        workbook.save(temporary)
        os.replace(temporary, output_path)
    finally:
        workbook.close()


def validate_result_xlsx(path: Path, expected: SimulationResult) -> dict:
    """Validate workbook structure and every written numerical value."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["温度", "水分浓度"]:
            raise ValueError(f"unexpected output worksheets: {workbook.sheetnames}")
        details = {}
        expected_fields = {
            "温度": expected.temperatures_c,
            "水分浓度": expected.moistures_kg_kg,
        }
        for sheet_name, expected_field in expected_fields.items():
            worksheet = workbook[sheet_name]
            if worksheet.max_row != 10801 or worksheet.max_column != 22:
                raise ValueError(
                    f"{sheet_name} has {worksheet.max_row}x{worksheet.max_column}, "
                    "expected 10801x22"
                )
            header = [worksheet.cell(1, column).value for column in range(2, 23)]
            expected_header = [0.1 * i for i in range(21)]
            if any(abs(float(a) - b) > 1.0e-12 for a, b in zip(header, expected_header)):
                raise ValueError(f"unexpected radius header in {sheet_name}")
            if worksheet["A2"].value != 1 or worksheet["A10801"].value != 10800:
                raise ValueError(f"unexpected time range in {sheet_name}")

            max_write_error = 0.0
            checked_values = 0
            for row_index, row in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True)
            ):
                if int(row[0]) != int(expected.times_s[row_index]):
                    raise ValueError(
                        f"unexpected time at {sheet_name}!A{row_index + 2}: {row[0]}"
                    )
                written = np.asarray(row[1:], dtype=float)
                wanted = np.round(expected_field[row_index], decimals=4)
                if written.shape != wanted.shape or not np.all(np.isfinite(written)):
                    raise ValueError(
                        f"invalid numerical row at {sheet_name}!{row_index + 2}"
                    )
                max_write_error = max(
                    max_write_error, float(np.max(np.abs(written - wanted)))
                )
                checked_values += written.size
            if max_write_error > 5.0e-13:
                raise ValueError(
                    f"written values differ from rounded solver output in {sheet_name}: "
                    f"max error {max_write_error:.3e}"
                )
            details[sheet_name] = {
                "rows": worksheet.max_row,
                "columns": worksheet.max_column,
                "dimension": worksheet.calculate_dimension(),
                "checked_numerical_values": checked_values,
                "max_error_against_rounded_solver_output": max_write_error,
            }
        return {"sheet_names": workbook.sheetnames, "worksheets": details}
    finally:
        workbook.close()


def select_rows(result: SimulationResult, requested_times_s: Sequence[int]) -> list[int]:
    time_to_index = {int(t): i for i, t in enumerate(result.times_s)}
    return [time_to_index[int(t)] for t in requested_times_s]


def markdown_table(
    times_h: Sequence[float], radii_cm: Sequence[float], values: np.ndarray
) -> str:
    header = "| 时间/h | " + " | ".join(f"{radius:g} cm" for radius in radii_cm) + " |"
    separator = "|---:" + "|---:" * len(radii_cm) + "|"
    lines = [header, separator]
    for time_h, row in zip(times_h, values):
        lines.append(
            "| "
            + f"{time_h:.1f}"
            + " | "
            + " | ".join(f"{value:.4f}" for value in row)
            + " |"
        )
    return "\n".join(lines)


def write_summary_markdown(path: Path, result: SimulationResult) -> None:
    times_h = [0.5, 1.0, 1.5, 2.0, 2.5, 3.0]
    times_s = [int(time_h * 3600) for time_h in times_h]
    radii_cm = [0.0, 0.5, 1.0, 1.5, 2.0]
    time_indices = select_rows(result, times_s)
    radius_indices = [
        int(np.where(np.isclose(result.radii_cm, radius_cm, atol=1.0e-12))[0][0])
        for radius_cm in radii_cm
    ]
    temperature = result.temperatures_c[np.ix_(time_indices, radius_indices)]
    moisture = result.moistures_kg_kg[np.ix_(time_indices, radius_indices)]
    text = f"""# 第二问数值结果

结果由 `solve_problem2.py` 使用守恒型径向有限体积法和自适应隐式 BDF 时间积分计算。表中数值按题意保留四位小数；内部计算未提前舍入。

## 表 3：3 小时内药材的温度（°C）

{markdown_table(times_h, radii_cm, temperature)}

## 表 4：3 小时内药材的水分浓度（kg/kg）

{markdown_table(times_h, radii_cm, moisture)}

## 输出说明

- `result2.xlsx`：题目模板格式，含“温度”和“水分浓度”两个工作表；
- `temperature_full_precision.csv`：每秒、每 0.1 cm 的高精度温度结果；
- `moisture_full_precision.csv`：每秒、每 0.1 cm 的高精度水分浓度结果；
- `validation_report.md`：输入、物性、收敛性和输出文件检查。
"""
    path.write_text(text, encoding="utf-8")


def format_float(value: float) -> str:
    if math.isfinite(value):
        return f"{value:.6e}"
    return "未能从最大范数差得到正的收敛阶"


def write_validation_markdown(
    path: Path,
    environment: EnvironmentData,
    fine: SimulationResult,
    convergence: dict | None,
    temporal: dict | None,
    input_path: Path,
    template_path: Path,
    output_validation: dict,
    run_configs: dict[str, NumericalConfig],
    parameters: PhysicalParameters,
) -> None:
    d = fine.diagnostics
    grid = make_radial_grid(
        parameters.radius_m,
        fine.config.radial_cells,
        fine.config.surface_layer_m,
        fine.config.surface_refinement_factor,
    )
    grid_steps = np.diff(grid)
    lines = [
        "# 第二问数值验证报告",
        "",
        "## 总体结论",
        "",
        "程序采用变物性热湿耦合模型、守恒型径向有限体积空间离散和 SciPy 自适应隐式 BDF 时间积分。主结果使用表面局部加密网格，并通过空间网格加密和同网格时间容差加严进行复核。",
        "",
        "## 输入与模型口径",
        "",
        f"- 附件 1：`{input_path}`；SHA-256：`{sha256_file(input_path)}`；",
        f"- 模板：`{template_path}`；SHA-256：`{sha256_file(template_path)}`；",
        f"- 环境数据行数：{environment.times_s.size}；覆盖 {environment.times_s[0]:g}–{environment.times_s[-1]:g} s；",
        "- 第二问计算区间：0–10800 s，未对附件 1 作超范围外推；",
        "- 环境数据在相邻 60 s 测点间采用分段线性插值；",
        "- 第二问从初始状态重新计算，全部时段统一采用附录 3 物性公式；",
        "- 题面未另给第二问的表面对流系数，程序明确沿用第一问的 $h=25\\,\\mathrm{W/(m^2\\cdot K)}$ 与 $h_m=8\\times10^{-7}\\,\\mathrm{m/s}$。",
        "",
        "## 主计算配置",
        "",
        f"- 径向区间数：{fine.config.radial_cells}；节点数：{fine.config.radial_cells + 1}；",
        f"- 最小/最大空间步长：{np.min(grid_steps):.8g} / {np.max(grid_steps):.8g} m；",
        f"- BDF 相对容差：{fine.config.relative_tolerance:.3e}；",
        f"- 温度/含水率绝对容差：{fine.config.temperature_absolute_tolerance:.3e} / {fine.config.moisture_absolute_tolerance:.3e}；",
        f"- 最大内部时间步：{fine.config.maximum_step_s:g} s；初始时间步：{fine.config.first_step_s:g} s；",
        f"- 运行时间：{d.runtime_s:.3f} s；函数计算 {d.function_evaluations} 次；稀疏 LU 分解 {d.lu_decompositions} 次。",
        "",
        "## 数值范围与物性范围",
        "",
        f"- 温度：{d.min_temperature_c:.10f}–{d.max_temperature_c:.10f} °C；",
        f"- 水分浓度：{d.min_moisture_kg_kg:.10f}–{d.max_moisture_kg_kg:.10f} kg/kg；",
        f"- 密度：{d.min_density_kg_m3:.8f}–{d.max_density_kg_m3:.8f} kg/m³；",
        f"- 比热容：{d.min_heat_capacity_j_kg_k:.8f}–{d.max_heat_capacity_j_kg_k:.8f} J/(kg·K)；",
        f"- 导热系数：{d.min_conductivity_w_m_k:.10f}–{d.max_conductivity_w_m_k:.10f} W/(m·K)；",
        f"- 水分扩散系数：{d.min_diffusivity_m2_s:.10e}–{d.max_diffusivity_m2_s:.10e} m²/s；",
        f"- 以每秒输出作梯形积分得到的最大水分守恒绝对残差：{d.max_one_second_moisture_balance_abs:.6e}；",
        f"- 对应最大相对残差：{d.max_one_second_moisture_balance_rel:.6e}。该值包含 1 s 输出采样的积分误差，不是 BDF 内部残差。",
        "",
    ]
    if convergence is not None:
        lines.extend(
            [
                "## 三层空间网格收敛",
                "",
                "三层网格分别为 190、380、760 个径向区间；每次加密均将两区网格步长减半。比较范围包括 1–10800 s、0–2 cm 的全部输出点。",
                "",
                "| 场变量 | 中网格—细网格最大差 | 观测阶 | 细网格 Richardson 估计误差 | 1.25 安全系数上界 | 最大差位置 |",
                "|---|---:|---:|---:|---:|---|",
            ]
        )
        for label, field in (("温度/°C", "temperature"), ("水分浓度/(kg/kg)", "moisture")):
            item = convergence[field]
            lines.append(
                f"| {label} | {item['max_abs_difference']:.6e} | "
                f"{item['observed_order_max_norm']:.4f} | "
                f"{format_float(item['estimated_fine_grid_max_error'])} | "
                f"{format_float(item['estimated_fine_grid_max_error_with_1p25_safety'])} | "
                f"$t={item['max_location_time_s']}\\,\\mathrm{{s}}$, "
                f"$r={item['max_location_radius_cm']:.1f}\\,\\mathrm{{cm}}$ |"
            )
        lines.append("")
    if temporal is not None:
        lines.extend(
            [
                "## 同网格时间积分复核",
                "",
                "在 760 区间网格上，将 BDF 相对/绝对容差进一步缩小 4 倍，并将最大内部时间步由 5 s 缩小为 2.5 s。",
                "",
                "| 场变量 | 主计算—加严计算最大差 | RMS 差 | 最大差位置 |",
                "|---|---:|---:|---|",
            ]
        )
        for label, field in (("温度/°C", "temperature"), ("水分浓度/(kg/kg)", "moisture")):
            item = temporal[field]
            lines.append(
                f"| {label} | {item['max_abs_difference']:.6e} | "
                f"{item['rms_difference']:.6e} | "
                f"$t={item['max_location_time_s']}\\,\\mathrm{{s}}$, "
                f"$r={item['max_location_radius_cm']:.1f}\\,\\mathrm{{cm}}$ |"
            )
        lines.append("")
    lines.extend(
        [
            "## Excel 文件检查",
            "",
            f"- 工作表：{', '.join(output_validation['sheet_names'])}；",
            f"- 温度工作表：{output_validation['worksheets']['温度']['dimension']}；",
            f"- 水分浓度工作表：{output_validation['worksheets']['水分浓度']['dimension']}；",
            "- 每个工作表均为 10801 行、22 列，时间为 1–10800 s，空间为 0–2.0 cm；",
            f"- 已逐格核验 {sum(item['checked_numerical_values'] for item in output_validation['worksheets'].values()):,} 个数值单元格，与内存中求解结果四舍五入至四位小数后的最大差为 "
            f"{max(item['max_error_against_rounded_solver_output'] for item in output_validation['worksheets'].values()):.3e}；",
            "- Excel 数值写入前四舍五入至四位小数，高精度结果另存 CSV。",
            "",
            "## 复现实验配置",
            "",
            "```json",
            json.dumps({name: asdict(config) for name, config in run_configs.items()}, ensure_ascii=False, indent=2),
            "```",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def run_self_tests() -> None:
    parameters = PhysicalParameters()
    moisture = np.asarray([2.55])
    temperature = np.asarray([28.0])
    density, heat_capacity, conductivity, diffusivity = parameters.evaluate_properties(
        moisture, temperature
    )
    if abs(float(density[0]) - 976.4) > 1.0e-12:
        raise AssertionError("density formula self-test failed")
    expected_cp = 1450.0 + 2736.0 * 2.55 / 3.55
    expected_k = 0.21 + 0.38 * 2.55 / 3.55
    expected_d = 2.4e-3 * math.exp(-0.45 / 2.55) * math.exp(-3850.0 / 301.15)
    if abs(float(heat_capacity[0]) - expected_cp) > 1.0e-10:
        raise AssertionError("heat-capacity formula self-test failed")
    if abs(float(conductivity[0]) - expected_k) > 1.0e-14:
        raise AssertionError("conductivity formula self-test failed")
    if abs(float(diffusivity[0]) - expected_d) > 1.0e-22:
        raise AssertionError("diffusivity formula self-test failed")

    grid = make_radial_grid(0.02, 190, 0.002, 10)
    weights = radial_storage_weights(grid)
    if abs(float(np.sum(weights)) - 0.02**2 / 2.0) > 1.0e-16:
        raise AssertionError("radial control-volume weights self-test failed")
    requested_output_indices(grid)

    equilibrium_environment = EnvironmentData(
        np.asarray([0.0, 1.0]),
        np.asarray([28.0, 28.0]),
        np.asarray([2.55, 2.55]),
    )
    system = CoupledRadialSystem(equilibrium_environment, parameters, grid)
    state = np.concatenate((np.full(grid.size, 28.0), np.full(grid.size, 2.55)))
    rate = system(0.5, state)
    if float(np.max(np.abs(rate))) > 1.0e-12:
        raise AssertionError("uniform equilibrium preservation self-test failed")

    sparsity = make_jacobian_sparsity(grid.size)
    if sparsity.shape != (2 * grid.size, 2 * grid.size) or sparsity.nnz == 0:
        raise AssertionError("Jacobian sparsity self-test failed")


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    project_root = Path(__file__).resolve().parent.parent
    output_directory = project_root / "result"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=project_root / "A题" / "附件" / "附件1.xlsx",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=project_root / "A题" / "附件" / "附件3" / "result2.xlsx",
    )
    parser.add_argument("--outdir", type=Path, default=output_directory)
    parser.add_argument("--radial-cells", type=int, default=760)
    parser.add_argument("--rtol", type=float, default=2.0e-9)
    parser.add_argument("--max-step", type=float, default=5.0)
    parser.add_argument("--skip-convergence", action="store_true")
    parser.add_argument("--skip-temporal-check", action="store_true")
    parser.add_argument("--self-test-only", action="store_true")
    parser.add_argument("--quiet", action="store_true")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_arguments(argv)
    run_self_tests()
    if args.self_test_only:
        print("self-tests passed")
        return 0

    args.outdir.mkdir(parents=True, exist_ok=True)
    parameters = PhysicalParameters()
    main_config = NumericalConfig(
        radial_cells=args.radial_cells,
        relative_tolerance=args.rtol,
        maximum_step_s=args.max_step,
    )
    environment = load_environment(args.input, main_config.end_time_s)

    label = None if args.quiet else "fine/main"
    fine = solve_model(environment, parameters, main_config, label)
    convergence = None
    run_configs: dict[str, NumericalConfig] = {"fine_main": main_config}

    if not args.skip_convergence:
        if args.radial_cells % 4 != 0:
            raise ValueError("radial-cells must be divisible by four for convergence runs")
        medium_config = replace(main_config, radial_cells=args.radial_cells // 2)
        coarse_config = replace(main_config, radial_cells=args.radial_cells // 4)
        medium = solve_model(
            environment,
            parameters,
            medium_config,
            None if args.quiet else "medium",
        )
        coarse = solve_model(
            environment,
            parameters,
            coarse_config,
            None if args.quiet else "coarse",
        )
        coarse_to_medium = compare_results(coarse, medium)
        medium_to_fine = compare_results(medium, fine)
        convergence = add_richardson_estimates(
            coarse_to_medium, medium_to_fine
        )
        run_configs["medium"] = medium_config
        run_configs["coarse"] = coarse_config

    temporal = None
    if not args.skip_temporal_check:
        tight_config = replace(
            main_config,
            relative_tolerance=main_config.relative_tolerance / 4.0,
            temperature_absolute_tolerance=(
                main_config.temperature_absolute_tolerance / 4.0
            ),
            moisture_absolute_tolerance=(
                main_config.moisture_absolute_tolerance / 4.0
            ),
            maximum_step_s=main_config.maximum_step_s / 2.0,
            first_step_s=main_config.first_step_s / 2.0,
        )
        tight = solve_model(
            environment,
            parameters,
            tight_config,
            None if args.quiet else "fine/tight-time",
        )
        temporal = compare_results(fine, tight)
        run_configs["fine_tight_time"] = tight_config

    result_xlsx = args.outdir / "result2.xlsx"
    write_result_xlsx(args.template, result_xlsx, fine)
    output_validation = validate_result_xlsx(result_xlsx, fine)
    write_full_precision_csv(
        args.outdir / "temperature_full_precision.csv",
        fine.times_s,
        fine.radii_cm,
        fine.temperatures_c,
    )
    write_full_precision_csv(
        args.outdir / "moisture_full_precision.csv",
        fine.times_s,
        fine.radii_cm,
        fine.moistures_kg_kg,
    )
    write_environment_csv(
        args.outdir / "environment_input_used.csv",
        environment,
        main_config.end_time_s,
    )
    write_summary_markdown(args.outdir / "summary_tables.md", fine)
    write_validation_markdown(
        args.outdir / "validation_report.md",
        environment,
        fine,
        convergence,
        temporal,
        args.input,
        args.template,
        output_validation,
        run_configs,
        parameters,
    )

    metadata = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "python_version": sys.version,
        "platform": platform.platform(),
        "numpy_version": np.__version__,
        "scipy_version": scipy.__version__,
        "input": {
            "path": str(args.input.resolve()),
            "sha256": sha256_file(args.input),
        },
        "template": {
            "path": str(args.template.resolve()),
            "sha256": sha256_file(args.template),
        },
        "physical_parameters": asdict(parameters),
        "run_configs": {name: asdict(config) for name, config in run_configs.items()},
        "main_diagnostics": asdict(fine.diagnostics),
        "spatial_convergence": convergence,
        "temporal_check": temporal,
        "output_validation": output_validation,
        "outputs": {
            "result2_xlsx_sha256": sha256_file(result_xlsx),
        },
    }
    (args.outdir / "run_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    if not args.quiet:
        print(f"wrote {result_xlsx}")
        print(f"wrote {args.outdir / 'summary_tables.md'}")
        print(f"wrote {args.outdir / 'validation_report.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
