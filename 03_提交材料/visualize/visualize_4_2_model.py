"""True-density, non-affine shrinking-cylinder comparison model.

The measured radius is imposed. Appendix-4 density is interpreted as the
actual wet bulk density. A fixed dry-mass coordinate enforces local dry-solid
mass conservation exactly. The otherwise unobserved axial length is inferred
from global dry-mass conservation, so radial material motion is non-affine.
"""

from __future__ import annotations

import math
import time
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from numba import njit
from openpyxl import load_workbook


@dataclass(frozen=True)
class NumericalConfig:
    radial_cells: int = 760
    time_step_s: float = 2.0
    output_interval_s: float = 60.0
    maximum_time_h: float = 96.0
    surface_layer: float = 0.1
    surface_refinement_factor: int = 10
    picard_tolerance: float = 1.0e-10
    picard_max_iterations: int = 40
    drying_threshold_kg_kg: float = 0.15

    def validate(self) -> None:
        if self.radial_cells < 38 or self.radial_cells % 19 != 0:
            raise ValueError("radial_cells must be a multiple of 19 and at least 38")
        if self.time_step_s <= 0.0 or self.output_interval_s <= 0.0:
            raise ValueError("time intervals must be positive")
        stride = self.output_interval_s / self.time_step_s
        if abs(stride - round(stride)) > 1.0e-12:
            raise ValueError("output interval must be divisible by the time step")
        if self.maximum_time_h <= 4.0:
            raise ValueError("maximum time must exceed the measured environment")
        if not 0.0 < self.surface_layer < 1.0:
            raise ValueError("surface_layer must lie in (0,1)")
        if self.surface_refinement_factor < 1:
            raise ValueError("surface_refinement_factor must be positive")
        if self.picard_tolerance <= 0.0 or self.picard_max_iterations < 1:
            raise ValueError("invalid Picard settings")
        if not 0.05 < self.drying_threshold_kg_kg < 2.55:
            raise ValueError("invalid drying threshold")


@dataclass
class SimulationResult:
    times_s: np.ndarray
    dry_mass_coordinates: np.ndarray
    temperatures_c: np.ndarray
    moistures_kg_kg: np.ndarray
    material_radii_m: np.ndarray
    measured_radii_m: np.ndarray
    inferred_lengths_m: np.ndarray
    surface_temperatures_c: np.ndarray
    surface_moistures_kg_kg: np.ndarray
    maximum_moistures_kg_kg: np.ndarray
    drying_crossing_s: float
    runtime_s: float
    diagnostics: dict[str, float | int | bool]
    config: NumericalConfig


def load_two_column_or_three_column_xlsx(
    path: Path, expected_columns: int
) -> np.ndarray:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = []
        for row in workbook.active.iter_rows(min_row=2, values_only=True):
            if not row or all(value is None for value in row[:expected_columns]):
                continue
            if any(value is None for value in row[:expected_columns]):
                raise ValueError(f"missing value in {path}")
            rows.append([float(value) for value in row[:expected_columns]])
    finally:
        workbook.close()
    values = np.asarray(rows, dtype=float)
    if values.ndim != 2 or values.shape[1] != expected_columns:
        raise ValueError(f"unexpected spreadsheet shape: {path}")
    if not np.all(np.isfinite(values)) or np.any(np.diff(values[:, 0]) <= 0.0):
        raise ValueError(f"invalid time series: {path}")
    return values


def load_environment(path: Path) -> np.ndarray:
    values = load_two_column_or_three_column_xlsx(path, 3)
    if values[0, 0] != 0.0 or np.any(values[:, 2] <= 0.0):
        raise ValueError("invalid environment data")
    return values


def load_radius(path: Path) -> np.ndarray:
    values = load_two_column_or_three_column_xlsx(path, 2)
    values[:, 1] *= 0.01
    if values[0, 0] != 0.0 or np.any(values[:, 1] <= 0.0):
        raise ValueError("invalid radius data")
    return values


def make_mass_faces(
    radial_cells: int, surface_layer: float, refinement_factor: int
) -> np.ndarray:
    """Two-zone grid in normalized cumulative dry-mass coordinate."""
    bulk_length = 1.0 - surface_layer
    effective_length = bulk_length + refinement_factor * surface_layer
    bulk_step = effective_length / radial_cells
    surface_step = bulk_step / refinement_factor
    bulk_cells = int(round(bulk_length / bulk_step))
    surface_cells = int(round(surface_layer / surface_step))
    if bulk_cells + surface_cells != radial_cells:
        raise ValueError("radial_cells incompatible with the two-zone grid")
    bulk = np.arange(bulk_cells + 1, dtype=float) * bulk_step
    surface = bulk_length + np.arange(1, surface_cells + 1) * surface_step
    faces = np.concatenate((bulk, surface))
    faces[0] = 0.0
    faces[-1] = 1.0
    if faces.size != radial_cells + 1 or np.any(np.diff(faces) <= 0.0):
        raise AssertionError("invalid dry-mass grid")
    return faces


def appendix4_properties(
    moisture: np.ndarray, temperature_c: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    moisture = np.asarray(moisture, dtype=float)
    temperature_c = np.asarray(temperature_c, dtype=float)
    wet_density = 760.0 + 90.0 * moisture
    dry_density = wet_density / (1.0 + moisture)
    heat_capacity = 1850.0 + 2150.0 * moisture / (1.0 + moisture)
    conductivity = 0.12 + 0.20 * moisture / (1.0 + moisture)
    diffusivity = 4.2e-4 * np.exp(
        -0.30 / moisture - 3850.0 / (temperature_c + 273.15)
    )
    return wet_density, dry_density, heat_capacity, conductivity, diffusivity


def conservative_geometry(
    moisture: np.ndarray,
    mass_faces: np.ndarray,
    measured_radius_m: float,
    initial_radius_m: float = 0.02,
    initial_length_m: float = 0.25,
) -> tuple[np.ndarray, np.ndarray, float, np.ndarray]:
    """Return non-affine cell geometry and inferred length.

    Piecewise-constant cell density is used. Each mass-coordinate cell keeps
    exactly its initial fraction of total dry-solid mass.
    """
    moisture = np.asarray(moisture, dtype=float)
    mass_faces = np.asarray(mass_faces, dtype=float)
    widths = np.diff(mass_faces)
    if moisture.shape != widths.shape:
        raise ValueError("moisture and dry-mass grid shapes differ")
    _, dry_density, _, _, _ = appendix4_properties(
        moisture, np.full_like(moisture, 28.0)
    )
    initial_dry_density = (760.0 + 90.0 * 2.55) / (1.0 + 2.55)
    specific_volume_integral = float(np.sum(widths / dry_density))
    length_m = (
        initial_length_m
        * initial_radius_m**2
        * initial_dry_density
        * specific_volume_integral
        / measured_radius_m**2
    )
    cumulative = np.concatenate(
        ([0.0], np.cumsum(widths / dry_density) / specific_volume_integral)
    )
    radius_faces_m = measured_radius_m * np.sqrt(cumulative)
    radius_centers_m = np.sqrt(
        0.5 * (radius_faces_m[:-1] ** 2 + radius_faces_m[1:] ** 2)
    )
    local_ratios = (
        length_m
        * dry_density
        * np.diff(radius_faces_m**2)
        / (initial_length_m * initial_radius_m**2 * initial_dry_density * widths)
    )
    return radius_faces_m, radius_centers_m, length_m, local_ratios


@njit(cache=True)
def _solve_tridiagonal(lower, diagonal, upper, rhs, cprime, dprime, out):
    n = diagonal.size
    pivot = diagonal[0]
    if n > 1:
        cprime[0] = upper[0] / pivot
    dprime[0] = rhs[0] / pivot
    for i in range(1, n):
        pivot = diagonal[i] - lower[i - 1] * cprime[i - 1]
        if i < n - 1:
            cprime[i] = upper[i] / pivot
        dprime[i] = (rhs[i] - lower[i - 1] * dprime[i - 1]) / pivot
    out[n - 1] = dprime[n - 1]
    for i in range(n - 2, -1, -1):
        out[i] = dprime[i] - cprime[i] * out[i + 1]


@njit(cache=True)
def _implicit_step(
    old,
    previous,
    has_previous,
    inverse_dt,
    storage,
    conductance,
    boundary_conductance,
    external_value,
    lower,
    diagonal,
    upper,
    rhs,
    cprime,
    dprime,
    out,
):
    n = old.size
    time_diagonal = 1.5 * inverse_dt if has_previous else inverse_dt
    for i in range(n - 1):
        lower[i] = -conductance[i]
        upper[i] = -conductance[i]
    for i in range(n):
        if has_previous:
            history = (4.0 * old[i] - previous[i]) * (0.5 * inverse_dt)
        else:
            history = old[i] * inverse_dt
        west = conductance[i - 1] if i > 0 else 0.0
        east = conductance[i] if i < n - 1 else 0.0
        diagonal[i] = storage[i] * time_diagonal + west + east
        rhs[i] = storage[i] * history
    diagonal[n - 1] += boundary_conductance
    rhs[n - 1] += boundary_conductance * external_value
    _solve_tridiagonal(lower, diagonal, upper, rhs, cprime, dprime, out)


@njit(cache=True)
def _properties(moisture, temperature, wet_density, dry_density, cp, k, diffusivity):
    for i in range(moisture.size):
        c = moisture[i]
        rho_w = 760.0 + 90.0 * c
        wet_density[i] = rho_w
        dry_density[i] = rho_w / (1.0 + c)
        cp[i] = 1850.0 + 2150.0 * c / (1.0 + c)
        k[i] = 0.12 + 0.20 * c / (1.0 + c)
        diffusivity[i] = 4.2e-4 * math.exp(
            -0.30 / c - 3850.0 / (temperature[i] + 273.15)
        )


@njit(cache=True)
def _geometry(
    dry_density,
    mass_widths,
    measured_radius,
    initial_radius,
    initial_length,
    initial_dry_density,
    radius_squared_faces,
    radius_centers,
):
    specific_volume = 0.0
    for i in range(dry_density.size):
        specific_volume += mass_widths[i] / dry_density[i]
    length = (
        initial_length
        * initial_radius
        * initial_radius
        * initial_dry_density
        * specific_volume
        / (measured_radius * measured_radius)
    )
    radius_squared_faces[0] = 0.0
    cumulative = 0.0
    for i in range(dry_density.size):
        cumulative += mass_widths[i] / dry_density[i]
        radius_squared_faces[i + 1] = (
            measured_radius * measured_radius * cumulative / specific_volume
        )
        radius_centers[i] = math.sqrt(
            0.5 * (radius_squared_faces[i] + radius_squared_faces[i + 1])
        )
    return length


@njit(cache=True)
def _conductances(
    mass_faces,
    mass_centers,
    dry_density,
    conductivity,
    diffusivity,
    radius_squared_faces,
    measured_radius,
    length,
    total_dry_mass,
    heat_transfer,
    mass_transfer,
    heat_conductance,
    moisture_conductance,
):
    n = dry_density.size
    scale = 4.0 * math.pi * math.pi * length * length / (total_dry_mass * total_dry_mass)
    for i in range(n - 1):
        face = mass_faces[i + 1]
        left_distance = face - mass_centers[i]
        right_distance = mass_centers[i + 1] - face
        heat_resistance = (
            left_distance / (dry_density[i] * conductivity[i])
            + right_distance / (dry_density[i + 1] * conductivity[i + 1])
        )
        moisture_resistance = (
            left_distance / (dry_density[i] * dry_density[i] * diffusivity[i])
            + right_distance
            / (dry_density[i + 1] * dry_density[i + 1] * diffusivity[i + 1])
        )
        heat_conductance[i] = scale * radius_squared_faces[i + 1] / heat_resistance
        moisture_conductance[i] = (
            scale * radius_squared_faces[i + 1] / moisture_resistance
        )

    half_cell = mass_faces[n] - mass_centers[n - 1]
    heat_diffusive = (
        scale
        * measured_radius
        * measured_radius
        * dry_density[n - 1]
        * conductivity[n - 1]
        / half_cell
    )
    moisture_diffusive = (
        scale
        * measured_radius
        * measured_radius
        * dry_density[n - 1]
        * dry_density[n - 1]
        * diffusivity[n - 1]
        / half_cell
    )
    heat_external = (
        2.0 * math.pi * length * measured_radius * heat_transfer / total_dry_mass
    )
    moisture_external = (
        2.0
        * math.pi
        * length
        * measured_radius
        * dry_density[n - 1]
        * mass_transfer
        / total_dry_mass
    )
    heat_boundary = heat_diffusive * heat_external / (heat_diffusive + heat_external)
    moisture_boundary = (
        moisture_diffusive
        * moisture_external
        / (moisture_diffusive + moisture_external)
    )
    return heat_boundary, moisture_boundary, heat_external, moisture_external


@njit(cache=True)
def _center_value(mass_centers, values):
    """Second-order reconstruction at the symmetry axis.

    A smooth axisymmetric field is even in r.  Since cumulative dry-mass
    coordinate is proportional to r**2 near the axis, the field is smooth and
    locally linear in mu.  Extrapolating the first two cell averages to mu=0
    avoids treating the first cell average as the pointwise center maximum.
    """
    if values.size < 2:
        return values[0]
    slope = (values[1] - values[0]) / (mass_centers[1] - mass_centers[0])
    return values[0] - mass_centers[0] * slope


@njit(cache=True)
def _sample_profile(
    mass_centers, values, surface_value, sample_coordinates, output
):
    n = mass_centers.size
    cursor = 0
    for j in range(sample_coordinates.size):
        target = sample_coordinates[j]
        if target <= mass_centers[0]:
            slope = (values[1] - values[0]) / (
                mass_centers[1] - mass_centers[0]
            )
            output[j] = values[0] + (target - mass_centers[0]) * slope
        elif target >= mass_centers[n - 1]:
            span = 1.0 - mass_centers[n - 1]
            fraction = (target - mass_centers[n - 1]) / span
            output[j] = values[n - 1] + fraction * (surface_value - values[n - 1])
        else:
            while cursor + 1 < n and mass_centers[cursor + 1] < target:
                cursor += 1
            fraction = (
                (target - mass_centers[cursor])
                / (mass_centers[cursor + 1] - mass_centers[cursor])
            )
            output[j] = values[cursor] + fraction * (
                values[cursor + 1] - values[cursor]
            )


@njit(cache=True)
def _sample_radii(
    mass_faces, radius_squared_faces, sample_coordinates, output
):
    n = mass_faces.size - 1
    cursor = 0
    for j in range(sample_coordinates.size):
        target = sample_coordinates[j]
        if target <= 0.0:
            output[j] = 0.0
        elif target >= 1.0:
            output[j] = math.sqrt(radius_squared_faces[n])
        else:
            while cursor + 1 < n and mass_faces[cursor + 1] < target:
                cursor += 1
            fraction = (
                (target - mass_faces[cursor])
                / (mass_faces[cursor + 1] - mass_faces[cursor])
            )
            radius_squared = radius_squared_faces[cursor] + fraction * (
                radius_squared_faces[cursor + 1] - radius_squared_faces[cursor]
            )
            output[j] = math.sqrt(radius_squared)


@njit(cache=True)
def _kernel(
    mass_faces,
    mass_centers,
    mass_widths,
    sample_coordinates,
    time_step,
    output_stride,
    external_temperatures,
    external_moistures,
    measured_radii,
    maximum_steps,
    picard_tolerance,
    picard_max_iterations,
    drying_threshold,
    progress_stride,
):
    n = mass_widths.size
    inverse_dt = 1.0 / time_step
    initial_radius = 0.02
    initial_length = 0.25
    initial_moisture = 2.55
    initial_wet_density = 760.0 + 90.0 * initial_moisture
    initial_dry_density = initial_wet_density / (1.0 + initial_moisture)
    total_dry_mass = math.pi * initial_length * initial_radius**2 * initial_dry_density

    temperature = np.full(n, 28.0)
    moisture = np.full(n, initial_moisture)
    previous_temperature = np.empty(n)
    previous_moisture = np.empty(n)
    guess_temperature = np.empty(n)
    guess_moisture = np.empty(n)
    candidate_temperature = np.empty(n)
    candidate_moisture = np.empty(n)
    wet_density = np.empty(n)
    dry_density = np.empty(n)
    heat_capacity = np.empty(n)
    conductivity = np.empty(n)
    diffusivity = np.empty(n)
    heat_storage = np.empty(n)
    moisture_storage = mass_widths.copy()
    radius_squared_faces = np.empty(n + 1)
    radius_centers = np.empty(n)
    heat_conductance = np.empty(n - 1)
    moisture_conductance = np.empty(n - 1)

    lower = np.empty(n - 1)
    diagonal = np.empty(n)
    upper = np.empty(n - 1)
    rhs = np.empty(n)
    cprime = np.empty(n - 1)
    dprime = np.empty(n)

    maximum_outputs = maximum_steps // output_stride + 3
    output_times = np.empty(maximum_outputs)
    output_temperatures = np.empty((maximum_outputs, sample_coordinates.size))
    output_moistures = np.empty((maximum_outputs, sample_coordinates.size))
    output_radii = np.empty((maximum_outputs, sample_coordinates.size))
    output_measured_radius = np.empty(maximum_outputs)
    output_length = np.empty(maximum_outputs)
    output_surface_temperature = np.empty(maximum_outputs)
    output_surface_moisture = np.empty(maximum_outputs)
    output_maximum_moisture = np.empty(maximum_outputs)
    sample_buffer = np.empty(sample_coordinates.size)
    output_count = 0

    max_picard = 0
    total_picard = 0
    max_water_balance_abs = 0.0
    max_water_balance_rel = 0.0
    max_local_dry_mass_error = 0.0
    min_length_ratio = 1.0e300
    max_length_ratio = -1.0e300
    max_radial_monotonicity_violation = 0.0
    has_previous = False
    previous_gap = initial_moisture - drying_threshold
    drying_time = -1.0

    # Initial output.
    _properties(
        moisture, temperature, wet_density, dry_density, heat_capacity,
        conductivity, diffusivity,
    )
    length = _geometry(
        dry_density, mass_widths, measured_radii[0], initial_radius,
        initial_length, initial_dry_density, radius_squared_faces, radius_centers,
    )
    heat_boundary, moisture_boundary, heat_external, moisture_external = _conductances(
        mass_faces, mass_centers, dry_density, conductivity, diffusivity,
        radius_squared_faces, measured_radii[0], length, total_dry_mass,
        25.0, 8.0e-7, heat_conductance, moisture_conductance,
    )
    heat_flux = heat_boundary * (external_temperatures[0] - temperature[n - 1])
    moisture_flux = moisture_boundary * (external_moistures[0] - moisture[n - 1])
    surface_temperature = external_temperatures[0] - heat_flux / heat_external
    surface_moisture = external_moistures[0] - moisture_flux / moisture_external
    output_times[0] = 0.0
    _sample_profile(
        mass_centers, temperature, surface_temperature, sample_coordinates,
        output_temperatures[0],
    )
    _sample_profile(
        mass_centers, moisture, surface_moisture, sample_coordinates,
        output_moistures[0],
    )
    _sample_radii(
        mass_faces, radius_squared_faces, sample_coordinates, output_radii[0]
    )
    output_measured_radius[0] = measured_radii[0]
    output_length[0] = length
    output_surface_temperature[0] = surface_temperature
    output_surface_moisture[0] = surface_moisture
    output_maximum_moisture[0] = initial_moisture
    output_count = 1

    for step in range(1, maximum_steps + 1):
        for i in range(n):
            guess_temperature[i] = temperature[i]
            guess_moisture[i] = moisture[i]

        converged = False
        iteration_used = 0
        heat_boundary = 0.0
        moisture_boundary = 0.0
        heat_external = 0.0
        moisture_external = 0.0
        length = initial_length
        for iteration in range(1, picard_max_iterations + 1):
            _properties(
                guess_moisture, guess_temperature, wet_density, dry_density,
                heat_capacity, conductivity, diffusivity,
            )
            length = _geometry(
                dry_density, mass_widths, measured_radii[step], initial_radius,
                initial_length, initial_dry_density, radius_squared_faces,
                radius_centers,
            )
            heat_boundary, _, heat_external, _ = _conductances(
                mass_faces, mass_centers, dry_density, conductivity, diffusivity,
                radius_squared_faces, measured_radii[step], length,
                total_dry_mass, 25.0, 8.0e-7, heat_conductance,
                moisture_conductance,
            )
            for i in range(n):
                heat_storage[i] = (
                    mass_widths[i]
                    * (1.0 + guess_moisture[i])
                    * heat_capacity[i]
                )
            _implicit_step(
                temperature, previous_temperature, has_previous, inverse_dt,
                heat_storage, heat_conductance, heat_boundary,
                external_temperatures[step], lower, diagonal, upper, rhs,
                cprime, dprime, candidate_temperature,
            )

            _properties(
                guess_moisture, candidate_temperature, wet_density, dry_density,
                heat_capacity, conductivity, diffusivity,
            )
            length = _geometry(
                dry_density, mass_widths, measured_radii[step], initial_radius,
                initial_length, initial_dry_density, radius_squared_faces,
                radius_centers,
            )
            _, moisture_boundary, _, moisture_external = _conductances(
                mass_faces, mass_centers, dry_density, conductivity, diffusivity,
                radius_squared_faces, measured_radii[step], length,
                total_dry_mass, 25.0, 8.0e-7, heat_conductance,
                moisture_conductance,
            )
            _implicit_step(
                moisture, previous_moisture, has_previous, inverse_dt,
                moisture_storage, moisture_conductance, moisture_boundary,
                external_moistures[step], lower, diagonal, upper, rhs,
                cprime, dprime, candidate_moisture,
            )

            maximum_update = 0.0
            for i in range(n):
                maximum_update = max(
                    maximum_update,
                    abs(candidate_temperature[i] - guess_temperature[i]) / 50.0,
                    abs(candidate_moisture[i] - guess_moisture[i]),
                )
                guess_temperature[i] = candidate_temperature[i]
                guess_moisture[i] = candidate_moisture[i]
            iteration_used = iteration
            if maximum_update < picard_tolerance:
                converged = True
                break
        if not converged:
            return (
                output_times[:output_count], output_temperatures[:output_count],
                output_moistures[:output_count], output_radii[:output_count],
                output_measured_radius[:output_count], output_length[:output_count],
                output_surface_temperature[:output_count],
                output_surface_moisture[:output_count],
                output_maximum_moisture[:output_count], -1.0,
                np.asarray([1.0, step]),
            )

        _properties(
            candidate_moisture, candidate_temperature, wet_density, dry_density,
            heat_capacity, conductivity, diffusivity,
        )
        length = _geometry(
            dry_density, mass_widths, measured_radii[step], initial_radius,
            initial_length, initial_dry_density, radius_squared_faces,
            radius_centers,
        )
        heat_boundary, moisture_boundary, heat_external, moisture_external = _conductances(
            mass_faces, mass_centers, dry_density, conductivity, diffusivity,
            radius_squared_faces, measured_radii[step], length, total_dry_mass,
            25.0, 8.0e-7, heat_conductance, moisture_conductance,
        )

        storage_rate = 0.0
        for i in range(n):
            if has_previous:
                time_term = (
                    1.5 * candidate_moisture[i]
                    - 2.0 * moisture[i]
                    + 0.5 * previous_moisture[i]
                ) * inverse_dt
            else:
                time_term = (candidate_moisture[i] - moisture[i]) * inverse_dt
            storage_rate += mass_widths[i] * time_term
        boundary_rate = moisture_boundary * (
            external_moistures[step] - candidate_moisture[n - 1]
        )
        balance_abs = abs(storage_rate - boundary_rate)
        balance_scale = max(abs(storage_rate), abs(boundary_rate), 1.0e-16)
        max_water_balance_abs = max(max_water_balance_abs, balance_abs)
        max_water_balance_rel = max(
            max_water_balance_rel, balance_abs / balance_scale
        )

        for i in range(n):
            local_ratio = (
                length
                * dry_density[i]
                * (radius_squared_faces[i + 1] - radius_squared_faces[i])
                / (
                    initial_length
                    * initial_radius
                    * initial_radius
                    * initial_dry_density
                    * mass_widths[i]
                )
            )
            max_local_dry_mass_error = max(
                max_local_dry_mass_error, abs(local_ratio - 1.0)
            )
            if i < n - 1:
                max_radial_monotonicity_violation = max(
                    max_radial_monotonicity_violation,
                    candidate_moisture[i + 1] - candidate_moisture[i],
                )
        length_ratio = length / initial_length
        min_length_ratio = min(min_length_ratio, length_ratio)
        max_length_ratio = max(max_length_ratio, length_ratio)
        max_picard = max(max_picard, iteration_used)
        total_picard += iteration_used

        heat_flux = heat_boundary * (
            external_temperatures[step] - candidate_temperature[n - 1]
        )
        moisture_flux = moisture_boundary * (
            external_moistures[step] - candidate_moisture[n - 1]
        )
        surface_temperature = (
            external_temperatures[step] - heat_flux / heat_external
        )
        surface_moisture = (
            external_moistures[step] - moisture_flux / moisture_external
        )
        center_moisture = _center_value(mass_centers, candidate_moisture)
        new_maximum = max(
            center_moisture, np.max(candidate_moisture), surface_moisture
        )
        new_gap = new_maximum - drying_threshold

        if new_gap < 0.0 <= previous_gap:
            fraction = previous_gap / (previous_gap - new_gap)
            fraction = min(1.0, max(0.0, fraction))
            drying_time = (step - 1 + fraction) * time_step
            for i in range(n):
                guess_temperature[i] = temperature[i] + fraction * (
                    candidate_temperature[i] - temperature[i]
                )
                guess_moisture[i] = moisture[i] + fraction * (
                    candidate_moisture[i] - moisture[i]
                )
            event_radius = measured_radii[step - 1] + fraction * (
                measured_radii[step] - measured_radii[step - 1]
            )
            _properties(
                guess_moisture, guess_temperature, wet_density, dry_density,
                heat_capacity, conductivity, diffusivity,
            )
            length = _geometry(
                dry_density, mass_widths, event_radius, initial_radius,
                initial_length, initial_dry_density, radius_squared_faces,
                radius_centers,
            )
            event_temperature_external = external_temperatures[step - 1] + fraction * (
                external_temperatures[step] - external_temperatures[step - 1]
            )
            event_moisture_external = external_moistures[step - 1] + fraction * (
                external_moistures[step] - external_moistures[step - 1]
            )
            heat_boundary, moisture_boundary, heat_external, moisture_external = _conductances(
                mass_faces, mass_centers, dry_density, conductivity, diffusivity,
                radius_squared_faces, event_radius, length, total_dry_mass,
                25.0, 8.0e-7, heat_conductance, moisture_conductance,
            )
            heat_flux = heat_boundary * (
                event_temperature_external - guess_temperature[n - 1]
            )
            moisture_flux = moisture_boundary * (
                event_moisture_external - guess_moisture[n - 1]
            )
            surface_temperature = event_temperature_external - heat_flux / heat_external
            surface_moisture = event_moisture_external - moisture_flux / moisture_external
            if abs(output_times[output_count - 1] - drying_time) > 1.0e-9:
                output_times[output_count] = drying_time
                _sample_profile(
                    mass_centers, guess_temperature, surface_temperature,
                    sample_coordinates, output_temperatures[output_count],
                )
                _sample_profile(
                    mass_centers, guess_moisture, surface_moisture,
                    sample_coordinates, output_moistures[output_count],
                )
                _sample_radii(
                    mass_faces, radius_squared_faces, sample_coordinates,
                    output_radii[output_count],
                )
                output_measured_radius[output_count] = event_radius
                output_length[output_count] = length
                output_surface_temperature[output_count] = surface_temperature
                output_surface_moisture[output_count] = surface_moisture
                output_maximum_moisture[output_count] = max(
                    _center_value(mass_centers, guess_moisture),
                    np.max(guess_moisture),
                    surface_moisture,
                )
                output_count += 1
            diagnostics = np.asarray(
                [
                    0.0, step, max_picard, total_picard,
                    max_water_balance_abs, max_water_balance_rel,
                    max_local_dry_mass_error, min_length_ratio,
                    max_length_ratio, max_radial_monotonicity_violation,
                ]
            )
            return (
                output_times[:output_count], output_temperatures[:output_count],
                output_moistures[:output_count], output_radii[:output_count],
                output_measured_radius[:output_count], output_length[:output_count],
                output_surface_temperature[:output_count],
                output_surface_moisture[:output_count],
                output_maximum_moisture[:output_count], drying_time, diagnostics,
            )

        if step % output_stride == 0:
            output_times[output_count] = step * time_step
            _sample_profile(
                mass_centers, candidate_temperature, surface_temperature,
                sample_coordinates, output_temperatures[output_count],
            )
            _sample_profile(
                mass_centers, candidate_moisture, surface_moisture,
                sample_coordinates, output_moistures[output_count],
            )
            _sample_radii(
                mass_faces, radius_squared_faces, sample_coordinates,
                output_radii[output_count],
            )
            output_measured_radius[output_count] = measured_radii[step]
            output_length[output_count] = length
            output_surface_temperature[output_count] = surface_temperature
            output_surface_moisture[output_count] = surface_moisture
            output_maximum_moisture[output_count] = new_maximum
            output_count += 1

        for i in range(n):
            previous_temperature[i] = temperature[i]
            temperature[i] = candidate_temperature[i]
            previous_moisture[i] = moisture[i]
            moisture[i] = candidate_moisture[i]
        has_previous = True
        previous_gap = new_gap
        if progress_stride > 0 and step % progress_stride == 0:
            print(
                "progress_h", step * time_step / 3600.0,
                "Cmax", new_maximum, "L_cm", 100.0 * length,
            )

    return (
        output_times[:output_count], output_temperatures[:output_count],
        output_moistures[:output_count], output_radii[:output_count],
        output_measured_radius[:output_count], output_length[:output_count],
        output_surface_temperature[:output_count],
        output_surface_moisture[:output_count],
        output_maximum_moisture[:output_count], -1.0,
        np.asarray([2.0, maximum_steps]),
    )


def solve_model(
    environment: np.ndarray,
    radius_data: np.ndarray,
    config: NumericalConfig,
    progress: bool = False,
) -> SimulationResult:
    config.validate()
    mass_faces = make_mass_faces(
        config.radial_cells,
        config.surface_layer,
        config.surface_refinement_factor,
    )
    mass_centers = 0.5 * (mass_faces[:-1] + mass_faces[1:])
    mass_widths = np.diff(mass_faces)
    sample_coordinates = np.linspace(0.0, 1.0, 21)
    maximum_steps = int(math.ceil(config.maximum_time_h * 3600.0 / config.time_step_s))
    times = np.arange(maximum_steps + 1, dtype=float) * config.time_step_s
    external_temperature = np.interp(times, environment[:, 0], environment[:, 1])
    external_moisture = np.interp(times, environment[:, 0], environment[:, 2])
    after_environment = times >= environment[-1, 0]
    external_temperature[after_environment] = 50.0
    external_moisture[after_environment] = 0.05
    measured_radius = np.interp(times, radius_data[:, 0], radius_data[:, 1])
    measured_radius[times >= radius_data[-1, 0]] = radius_data[-1, 1]

    started = time.perf_counter()
    values = _kernel(
        mass_faces,
        mass_centers,
        mass_widths,
        sample_coordinates,
        config.time_step_s,
        int(round(config.output_interval_s / config.time_step_s)),
        external_temperature,
        external_moisture,
        measured_radius,
        maximum_steps,
        config.picard_tolerance,
        config.picard_max_iterations,
        config.drying_threshold_kg_kg,
        int(round(21600.0 / config.time_step_s)) if progress else 0,
    )
    (
        output_times,
        output_temperatures,
        output_moistures,
        output_radii,
        output_measured_radius,
        output_length,
        output_surface_temperature,
        output_surface_moisture,
        output_maximum_moisture,
        drying_time,
        diagnostic_values,
    ) = values
    if drying_time < 0.0:
        reason = "Picard iteration failed" if diagnostic_values[0] == 1.0 else "threshold not reached"
        raise RuntimeError(
            f"{reason}; diagnostic step={int(diagnostic_values[1])}"
        )
    diagnostics: dict[str, float | int | bool] = {
        "steps": int(diagnostic_values[1]),
        "max_picard_iterations": int(diagnostic_values[2]),
        "mean_picard_iterations": float(diagnostic_values[3] / diagnostic_values[1]),
        "max_water_balance_abs": float(diagnostic_values[4]),
        "max_water_balance_rel": float(diagnostic_values[5]),
        "max_local_dry_mass_relative_error": float(diagnostic_values[6]),
        "minimum_length_ratio": float(diagnostic_values[7]),
        "maximum_length_ratio": float(diagnostic_values[8]),
        "max_radial_monotonicity_violation": float(diagnostic_values[9]),
        "maximum_moisture_at_center": bool(
            np.all(
                output_moistures[:, 0]
                >= np.max(output_moistures[:, 1:], axis=1) - 1.0e-10
            )
        ),
    }
    return SimulationResult(
        times_s=output_times,
        dry_mass_coordinates=sample_coordinates,
        temperatures_c=output_temperatures,
        moistures_kg_kg=output_moistures,
        material_radii_m=output_radii,
        measured_radii_m=output_measured_radius,
        inferred_lengths_m=output_length,
        surface_temperatures_c=output_surface_temperature,
        surface_moistures_kg_kg=output_surface_moisture,
        maximum_moistures_kg_kg=output_maximum_moisture,
        drying_crossing_s=float(drying_time),
        runtime_s=time.perf_counter() - started,
        diagnostics=diagnostics,
        config=config,
    )
