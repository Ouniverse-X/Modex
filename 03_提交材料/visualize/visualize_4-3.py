#!/usr/bin/env python3
"""Run, archive, compare and visualize the Q4-3 free-radius experiment."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import platform
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
from visualize_4_3_model import (
    NumericalConfig,
    SimulationResult,
    load_environment,
    load_radius,
    solve_model,
)

ENVIRONMENT_PATH = ROOT / "A题" / "附件" / "附件1.xlsx"
RADIUS_PATH = ROOT / "A题" / "附件" / "附件2.xlsx"
SOLVER_PATH = HERE / "visualize_4_3_model.py"
CORE_PATH = HERE / "visualize_4_2_model.py"
DEFAULT_OUTPUT = HERE / "experiments" / "fixed_length_free_radius_20260911_final"

CASES: dict[str, NumericalConfig] = {
    "space_n760_dt1": NumericalConfig(
        radial_cells=760, time_step_s=1.0, output_interval_s=60.0,
        maximum_time_h=72.0,
    ),
    "space_n1520_dt1": NumericalConfig(
        radial_cells=1520, time_step_s=1.0, output_interval_s=60.0,
        maximum_time_h=72.0,
    ),
    "reference_n3040_dt1": NumericalConfig(
        radial_cells=3040, time_step_s=1.0, output_interval_s=60.0,
        maximum_time_h=72.0,
    ),
    "time_n3040_dt2": NumericalConfig(
        radial_cells=3040, time_step_s=2.0, output_interval_s=60.0,
        maximum_time_h=72.0,
    ),
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_default(value: Any) -> Any:
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, Path):
        return str(value)
    raise TypeError(f"cannot serialize {type(value)!r}")


def save_result(
    result: SimulationResult, case_dir: Path, case_name: str
) -> dict[str, Any]:
    case_dir.mkdir(parents=True, exist_ok=False)
    np.savez_compressed(
        case_dir / "result_full_precision.npz",
        times_s=result.times_s,
        dry_mass_coordinates=result.dry_mass_coordinates,
        temperatures_c=result.temperatures_c,
        moistures_kg_kg=result.moistures_kg_kg,
        material_radii_m=result.material_radii_m,
        predicted_radii_m=result.predicted_radii_m,
        surface_temperatures_c=result.surface_temperatures_c,
        surface_moistures_kg_kg=result.surface_moistures_kg_kg,
        maximum_moistures_kg_kg=result.maximum_moistures_kg_kg,
    )
    metadata: dict[str, Any] = {
        "case": case_name,
        "config": asdict(result.config),
        "drying_crossing_s": result.drying_crossing_s,
        "drying_crossing_h": result.drying_crossing_s / 3600.0,
        "runtime_s": result.runtime_s,
        "diagnostics": result.diagnostics,
        "output_rows": int(result.times_s.size),
        "predicted_radius_initial_cm": float(100.0 * result.predicted_radii_m[0]),
        "predicted_radius_final_cm": float(100.0 * result.predicted_radii_m[-1]),
        "predicted_radius_min_cm": float(100.0 * np.min(result.predicted_radii_m)),
        "predicted_radius_max_cm": float(100.0 * np.max(result.predicted_radii_m)),
    }
    (case_dir / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, default=_json_default) + "\n",
        encoding="utf-8",
    )
    return metadata


def load_archived_result(case_dir: Path) -> dict[str, Any]:
    with np.load(case_dir / "result_full_precision.npz") as archive:
        result = {name: archive[name] for name in archive.files}
    result["metadata"] = json.loads(
        (case_dir / "metadata.json").read_text(encoding="utf-8")
    )
    return result


def run_case(
    case_name: str, config_dict: dict[str, Any], output_dir: str
) -> dict[str, Any]:
    environment = load_environment(ENVIRONMENT_PATH)
    result = solve_model(environment, NumericalConfig(**config_dict), progress=False)
    return save_result(result, Path(output_dir) / case_name, case_name)


def compare_cases(coarse: dict[str, Any], fine: dict[str, Any]) -> dict[str, float]:
    if not np.array_equal(coarse["times_s"], fine["times_s"]):
        raise ValueError("archived output times do not match")
    radius_difference_cm = 100.0 * np.abs(
        coarse["predicted_radii_m"] - fine["predicted_radii_m"]
    )
    return {
        "radius_max_abs_cm": float(np.max(radius_difference_cm)),
        "radius_final_abs_cm": float(radius_difference_cm[-1]),
        "moisture_max_abs": float(
            np.max(np.abs(coarse["moistures_kg_kg"] - fine["moistures_kg_kg"]))
        ),
        "temperature_max_abs_c": float(
            np.max(np.abs(coarse["temperatures_c"] - fine["temperatures_c"]))
        ),
        "crossing_difference_s": float(
            abs(
                coarse["metadata"]["drying_crossing_s"]
                - fine["metadata"]["drying_crossing_s"]
            )
        ),
    }


def observed_order(error_coarse_medium: float, error_medium_fine: float) -> float:
    if error_coarse_medium <= 0.0 or error_medium_fine <= 0.0:
        return float("nan")
    return math.log(error_coarse_medium / error_medium_fine, 2.0)


def attachment_comparison(reference: dict[str, Any]) -> tuple[np.ndarray, dict[str, float]]:
    observed = load_radius(RADIUS_PATH)
    predicted = np.interp(
        observed[:, 0], reference["times_s"], reference["predicted_radii_m"]
    )
    residual_cm = 100.0 * (predicted - observed[:, 1])
    absolute_cm = np.abs(residual_cm)
    relative_percent = 100.0 * (predicted - observed[:, 1]) / observed[:, 1]
    sse = float(np.sum((predicted - observed[:, 1]) ** 2))
    sst = float(np.sum((observed[:, 1] - np.mean(observed[:, 1])) ** 2))
    maximum_index = int(np.argmax(absolute_cm))
    table = np.column_stack(
        (
            observed[:, 0],
            observed[:, 0] / 3600.0,
            100.0 * observed[:, 1],
            100.0 * predicted,
            residual_cm,
            absolute_cm,
            relative_percent,
        )
    )
    metrics = {
        "sample_count": int(observed.shape[0]),
        "bias_cm": float(np.mean(residual_cm)),
        "mae_cm": float(np.mean(absolute_cm)),
        "rmse_cm": float(np.sqrt(np.mean(residual_cm**2))),
        "mape_percent": float(np.mean(np.abs(relative_percent))),
        "maximum_absolute_error_cm": float(absolute_cm[maximum_index]),
        "maximum_absolute_error_time_h": float(observed[maximum_index, 0] / 3600.0),
        "r_squared": float(1.0 - sse / sst),
        "observed_final_radius_cm": float(100.0 * observed[-1, 1]),
        "predicted_final_radius_cm": float(100.0 * predicted[-1]),
        "final_residual_cm": float(residual_cm[-1]),
        "all_noninitial_residuals_positive": bool(np.all(residual_cm[1:] > 0.0)),
    }
    return table, metrics


def style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        width = min(
            34,
            max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2),
        )
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def make_workbook(
    output_path: Path,
    reference: dict[str, Any],
    comparison_table: np.ndarray,
    metadata: dict[str, Any],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "结果摘要"
    summary.append(["指标", "数值", "单位/说明"])
    ref_meta = reference["metadata"]
    fit = metadata["attachment2_comparison"]
    rows = [
        ("模型", "真实湿密度 + 干质量守恒 + 固定长度", "附件2不参与正向计算"),
        ("固定长度", 25.0, "cm"),
        ("参考空间网格", ref_meta["config"]["radial_cells"], "累计干质量单元"),
        ("参考时间步", ref_meta["config"]["time_step_s"], "s"),
        ("输出间隔", ref_meta["config"]["output_interval_s"], "s"),
        ("预测初始半径", ref_meta["predicted_radius_initial_cm"], "cm"),
        ("预测72 h半径", fit["predicted_final_radius_cm"], "cm"),
        ("实测72 h半径", fit["observed_final_radius_cm"], "cm"),
        ("72 h残差（预测-实测）", fit["final_residual_cm"], "cm"),
        ("平均偏差（预测-实测）", fit["bias_cm"], "cm"),
        ("MAE", fit["mae_cm"], "cm"),
        ("RMSE", fit["rmse_cm"], "cm"),
        ("MAPE", fit["mape_percent"], "%"),
        ("最大绝对误差", fit["maximum_absolute_error_cm"], "cm"),
        ("最大绝对误差发生时刻", fit["maximum_absolute_error_time_h"], "h"),
        ("R²", fit["r_squared"], "仅作描述性指标"),
        ("Cmax=0.15连续临界时刻", ref_meta["drying_crossing_h"], "h"),
        (
            "逐环干质量最大相对误差",
            ref_meta["diagnostics"]["max_local_dry_mass_relative_error"],
            "-",
        ),
        (
            "水分收支最大绝对残差",
            ref_meta["diagnostics"]["max_water_balance_abs"],
            "s^-1",
        ),
        (
            "半径数值离散保守差值界",
            metadata["numerical_accuracy"]["radius_conservative_difference_bound_cm"],
            "cm（加密差值估计）",
        ),
    ]
    for row in rows:
        summary.append(row)
    style_sheet(summary)

    comparison = workbook.create_sheet("附件2逐点对比")
    comparison.append(
        [
            "time_s", "time_h", "observed_radius_cm", "predicted_radius_cm",
            "residual_pred_minus_obs_cm", "absolute_error_cm",
            "relative_error_percent",
        ]
    )
    for row in comparison_table:
        comparison.append([float(value) for value in row])
    style_sheet(comparison)

    history = workbook.create_sheet("参考解时序")
    history.append(
        [
            "time_s", "time_h", "predicted_radius_cm", "radius_ratio",
            "maximum_moisture_kg_kg", "center_moisture_kg_kg",
            "surface_moisture_kg_kg", "center_temperature_C",
            "surface_temperature_C",
        ]
    )
    for i, time_s in enumerate(reference["times_s"]):
        history.append(
            [
                float(time_s),
                float(time_s / 3600.0),
                float(100.0 * reference["predicted_radii_m"][i]),
                float(reference["predicted_radii_m"][i] / 0.02),
                float(reference["maximum_moistures_kg_kg"][i]),
                float(reference["moistures_kg_kg"][i, 0]),
                float(reference["surface_moistures_kg_kg"][i]),
                float(reference["temperatures_c"][i, 0]),
                float(reference["surface_temperatures_c"][i]),
            ]
        )
    style_sheet(history)

    profiles = workbook.create_sheet("代表时刻剖面")
    profiles.append(
        [
            "time_s", "time_h", "dry_mass_coordinate", "material_radius_cm",
            "temperature_C", "moisture_kg_kg",
        ]
    )
    for target_h in (0.0, 4.0, 12.0, 24.0, 36.0, 48.0, 60.0, 72.0):
        i = int(np.argmin(np.abs(reference["times_s"] / 3600.0 - target_h)))
        for j, coordinate in enumerate(reference["dry_mass_coordinates"]):
            profiles.append(
                [
                    float(reference["times_s"][i]),
                    float(reference["times_s"][i] / 3600.0),
                    float(coordinate),
                    float(100.0 * reference["material_radii_m"][i, j]),
                    float(reference["temperatures_c"][i, j]),
                    float(reference["moistures_kg_kg"][i, j]),
                ]
            )
    style_sheet(profiles)

    convergence = workbook.create_sheet("数值收敛性")
    convergence.append(["比较", "指标", "数值"])
    for name, values in metadata["convergence"]["comparisons"].items():
        for key, value in values.items():
            convergence.append([name, key, value])
    for key, value in metadata["convergence"]["orders"].items():
        convergence.append(["观测阶", key, value])
    for key, value in metadata["numerical_accuracy"].items():
        convergence.append(["精度结论", key, value])
    style_sheet(convergence)

    environment = workbook.create_sheet("附件1输入")
    environment.append(["time_s", "temperature_C", "ambient_moisture_kg_kg"])
    for row in load_environment(ENVIRONMENT_PATH):
        environment.append([float(value) for value in row])
    style_sheet(environment)

    radius = workbook.create_sheet("附件2输入")
    radius.append(["time_s", "observed_radius_cm"])
    for time_s, radius_m in load_radius(RADIUS_PATH):
        radius.append([float(time_s), float(100.0 * radius_m)])
    style_sheet(radius)

    workbook.save(output_path)


def configure_matplotlib() -> None:
    regular_font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    bold_font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")
    for font_path in (regular_font, bold_font):
        if font_path.exists():
            font_manager.fontManager.addfont(str(font_path))
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Noto Sans CJK JP", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "font.size": 11,
            "axes.labelsize": 11,
            "axes.titlesize": 13,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )


def make_figures(
    output_dir: Path,
    reference: dict[str, Any],
    comparison_table: np.ndarray,
    fit: dict[str, Any],
) -> None:
    figure_dir = output_dir / "figures"
    figure_dir.mkdir(exist_ok=True)
    configure_matplotlib()
    blue = "#2B6CB0"
    orange = "#DD6B20"
    gray = "#4A5568"
    chinese_font = font_manager.FontProperties(
        fname="/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
    )

    fig, (axis, residual_axis) = plt.subplots(
        2, 1, figsize=(9.2, 7.2), sharex=True,
        gridspec_kw={"height_ratios": [3.0, 1.15], "hspace": 0.08},
    )
    axis.plot(
        reference["times_s"] / 3600.0,
        100.0 * reference["predicted_radii_m"],
        color=orange,
        linewidth=2.3,
        label="检验模型预测",
        zorder=2,
    )
    axis.scatter(
        comparison_table[:, 1],
        comparison_table[:, 2],
        s=22,
        facecolors="white",
        edgecolors=blue,
        linewidths=1.1,
        label="附件2数据",
        zorder=3,
    )
    maximum_difference_index = int(np.argmax(np.abs(comparison_table[:, 4])))
    maximum_difference_time_h = comparison_table[maximum_difference_index, 1]
    observed_radius_cm = comparison_table[maximum_difference_index, 2]
    predicted_radius_cm = comparison_table[maximum_difference_index, 3]
    axis.annotate(
        "",
        xy=(maximum_difference_time_h, predicted_radius_cm),
        xytext=(maximum_difference_time_h, observed_radius_cm),
        arrowprops={
            "arrowstyle": "<->",
            "color": "#C53030",
            "linewidth": 2.0,
            "mutation_scale": 14,
        },
        zorder=5,
    )
    axis.set_ylabel(r"$r$ / cm", fontproperties=chinese_font)
    axis.grid(axis="y", color="#CBD5E0", linewidth=0.7, alpha=0.65)
    axis.legend(frameon=False, loc="upper right")
    residual_axis.axhline(0.0, color=gray, linewidth=1.0)
    residual_axis.plot(
        comparison_table[:, 1], comparison_table[:, 4],
        color=orange, linewidth=1.8,
    )
    residual_axis.fill_between(
        comparison_table[:, 1], 0.0, comparison_table[:, 4],
        color=orange, alpha=0.14,
    )
    residual_axis.set_xlabel("时间 / h")
    residual_axis.set_ylabel("预测值与实测值之差\n/ cm", fontproperties=chinese_font)
    residual_axis.set_xlim(0.0, 72.0)
    residual_axis.set_xticks(np.arange(0.0, 73.0, 12.0))
    residual_axis.grid(axis="y", color="#CBD5E0", linewidth=0.7, alpha=0.65)
    fig.align_ylabels((axis, residual_axis))
    for suffix in ("png", "svg", "pdf"):
        kwargs = {"dpi": 400} if suffix == "png" else {}
        fig.savefig(
            figure_dir / f"radius_attachment2_comparison.{suffix}",
            bbox_inches="tight",
            **kwargs,
        )
    plt.close(fig)

    times_h = reference["times_s"] / 3600.0
    fig, (radius_axis, moisture_axis) = plt.subplots(
        2, 1, figsize=(9.2, 7.0), sharex=True, constrained_layout=True
    )
    radius_axis.plot(
        times_h, 100.0 * reference["predicted_radii_m"],
        color=orange, linewidth=2.2,
    )
    radius_axis.set_ylabel("预测外半径 / cm")
    radius_axis.set_title("自由边界与内部水分场的耦合演化")
    radius_axis.grid(axis="y", color="#CBD5E0", linewidth=0.7, alpha=0.65)
    moisture_axis.plot(
        times_h, reference["maximum_moistures_kg_kg"],
        color=blue, linewidth=2.1, label="最大含水率",
    )
    moisture_axis.plot(
        times_h, reference["surface_moistures_kg_kg"],
        color=gray, linewidth=1.7, linestyle="--", label="表面含水率",
    )
    moisture_axis.axhline(0.15, color="#C53030", linewidth=1.1, linestyle=":", label="0.15阈值")
    moisture_axis.set_xlabel("时间 / h")
    moisture_axis.set_ylabel(r"干基含水率 / $\mathrm{(kg\,kg^{-1})}$")
    moisture_axis.set_xlim(0.0, 72.0)
    moisture_axis.set_xticks(np.arange(0.0, 73.0, 12.0))
    moisture_axis.grid(axis="y", color="#CBD5E0", linewidth=0.7, alpha=0.65)
    moisture_axis.legend(frameon=False, ncol=3, loc="upper right")
    for suffix in ("png", "svg"):
        kwargs = {"dpi": 400} if suffix == "png" else {}
        fig.savefig(
            figure_dir / f"radius_moisture_coupling.{suffix}",
            bbox_inches="tight",
            **kwargs,
        )
    plt.close(fig)


def write_report(output_dir: Path, metadata: dict[str, Any]) -> None:
    ref = metadata["cases"]["reference_n3040_dt1"]
    fit = metadata["attachment2_comparison"]
    conv = metadata["convergence"]
    accuracy = metadata["numerical_accuracy"]
    lines = [
        "# Q4-3 固定长度自由半径模型实验报告",
        "",
        "## 结论",
        "",
        "在不向模型输入附件2半径的前提下，本模型能够由真实湿密度、固定长度和逐环干质量守恒唯一预测外半径。"
        "但预测曲线除初始点外全部高于附件2实测值，呈系统性正偏，故该闭合方案在当前参数下不能充分解释实测收缩。",
        "",
        f"- 参考离散：$N={ref['config']['radial_cells']}$，$\\Delta t={ref['config']['time_step_s']:.0f}\\,\\mathrm{{s}}$，计算至72 h；",
        f"- 预测半径：{ref['predicted_radius_initial_cm']:.12f} cm（0 h）降至 {fit['predicted_final_radius_cm']:.12f} cm（72 h）；",
        f"- 实测72 h半径：{fit['observed_final_radius_cm']:.12f} cm，末点高估 {fit['final_residual_cm']:.12f} cm；",
        f"- 全部145个实测时刻：MAE={fit['mae_cm']:.12f} cm，RMSE={fit['rmse_cm']:.12f} cm，MAPE={fit['mape_percent']:.8f}%；",
        f"- 最大绝对误差={fit['maximum_absolute_error_cm']:.12f} cm，发生于 {fit['maximum_absolute_error_time_h']:.6f} h；",
        f"- 描述性 $R^2={fit['r_squared']:.10f}$；除初始点外残差是否全为正：{fit['all_noninitial_residuals_positive']}。",
        "",
        "## 守恒与数值自检",
        "",
        f"- 逐环干质量最大相对误差：{ref['diagnostics']['max_local_dry_mass_relative_error']:.3e}；",
        f"- 水分收支最大绝对残差：{ref['diagnostics']['max_water_balance_abs']:.3e} s⁻¹；",
        f"- 单步半径最大反向增加：{100.0 * ref['diagnostics']['max_single_step_radius_increase_m']:.3e} cm；",
        f"- $\\max C=0.15$ 连续临界时刻：{ref['drying_crossing_s']:.9f} s = {ref['drying_crossing_h']:.12f} h。",
        "",
        "## 网格与时间步收敛",
        "",
        "| 比较 | 半径全程最大差/cm | 72 h半径差/cm | 含水率最大差 | 临界时刻差/s |",
        "|---|---:|---:|---:|---:|",
    ]
    for name, item in conv["comparisons"].items():
        lines.append(
            f"| {name} | {item['radius_max_abs_cm']:.12e} | "
            f"{item['radius_final_abs_cm']:.12e} | {item['moisture_max_abs']:.12e} | "
            f"{item['crossing_difference_s']:.12e} |"
        )
    lines.extend(
        [
            "",
            f"半径空间观测阶：{conv['orders']['space_radius_max']:.8f}。"
            f"最细空间差与时间差之和乘1.5得到半径保守差值界 {accuracy['radius_conservative_difference_bound_cm']:.12e} cm；"
            f"该界相对于小数点后四位的半单位 {accuracy['four_decimal_half_unit_cm']:.1e} cm，"
            f"通过检查：{accuracy['passes_four_decimal_radius_guard']}。",
            "",
            "这里的误差界是网格加密差值形成的保守工程估计，不是数学意义上的严格后验误差定理。"
            "附件2的实验误差没有给出，因此上述MAE/RMSE是模型—数据偏差，不能直接解释为纯模型误差。",
        ]
    )
    (output_dir / "结果报告.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def summarize(output_dir: Path) -> dict[str, Any]:
    results = {name: load_archived_result(output_dir / name) for name in CASES}
    comparisons = {
        "space_760_vs_1520": compare_cases(
            results["space_n760_dt1"], results["space_n1520_dt1"]
        ),
        "space_1520_vs_3040": compare_cases(
            results["space_n1520_dt1"], results["reference_n3040_dt1"]
        ),
        "time_2s_vs_1s_at_n3040": compare_cases(
            results["time_n3040_dt2"], results["reference_n3040_dt1"]
        ),
    }
    orders = {
        "space_radius_max": observed_order(
            comparisons["space_760_vs_1520"]["radius_max_abs_cm"],
            comparisons["space_1520_vs_3040"]["radius_max_abs_cm"],
        ),
        "space_moisture_max": observed_order(
            comparisons["space_760_vs_1520"]["moisture_max_abs"],
            comparisons["space_1520_vs_3040"]["moisture_max_abs"],
        ),
        "space_crossing": observed_order(
            comparisons["space_760_vs_1520"]["crossing_difference_s"],
            comparisons["space_1520_vs_3040"]["crossing_difference_s"],
        ),
    }
    convergence = {"comparisons": comparisons, "orders": orders}
    comparison_table, attachment_metrics = attachment_comparison(
        results["reference_n3040_dt1"]
    )
    radius_bound = 1.5 * (
        comparisons["space_1520_vs_3040"]["radius_max_abs_cm"]
        + comparisons["time_2s_vs_1s_at_n3040"]["radius_max_abs_cm"]
    )
    accuracy = {
        "radius_conservative_difference_bound_cm": radius_bound,
        "four_decimal_half_unit_cm": 0.00005,
        "passes_four_decimal_radius_guard": bool(radius_bound < 0.00005),
        "qualification": "conservative grid-difference estimate, not a rigorous a posteriori theorem",
    }
    metadata: dict[str, Any] = {
        "model": "true wet density + local dry-mass conservation + fixed length + predicted free radius",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "independent_validation": "Attachment 2 radius is excluded from forward solve and used only for comparison",
        "inputs": {
            "environment": {
                "path": str(ENVIRONMENT_PATH),
                "sha256": sha256_file(ENVIRONMENT_PATH),
            },
            "validation_radius": {
                "path": str(RADIUS_PATH),
                "sha256": sha256_file(RADIUS_PATH),
            },
            "solver": {"path": str(SOLVER_PATH), "sha256": sha256_file(SOLVER_PATH)},
            "shared_numerical_core": {
                "path": str(CORE_PATH),
                "sha256": sha256_file(CORE_PATH),
            },
            "runner": {
                "path": str(Path(__file__).resolve()),
                "sha256": sha256_file(Path(__file__).resolve()),
            },
        },
        "environment": {
            "python": sys.version,
            "platform": platform.platform(),
            "cpu_count": os.cpu_count(),
        },
        "cases": {name: result["metadata"] for name, result in results.items()},
        "convergence": convergence,
        "numerical_accuracy": accuracy,
        "attachment2_comparison": attachment_metrics,
    }
    (output_dir / "run_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, default=_json_default) + "\n",
        encoding="utf-8",
    )
    make_workbook(
        output_dir / "result4-3.xlsx",
        results["reference_n3040_dt1"],
        comparison_table,
        metadata,
    )
    make_figures(
        output_dir,
        results["reference_n3040_dt1"],
        comparison_table,
        attachment_metrics,
    )
    write_report(output_dir, metadata)
    return metadata


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--summarize-only", action="store_true")
    args = parser.parse_args()
    output_dir = args.output.resolve()
    if output_dir.exists() and not args.resume and not args.summarize_only:
        raise FileExistsError(f"output already exists: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=args.resume or args.summarize_only)

    if not args.summarize_only:
        pending = {
            name: config
            for name, config in CASES.items()
            if not (output_dir / name / "result_full_precision.npz").exists()
        }
        if pending:
            with ProcessPoolExecutor(max_workers=min(args.workers, len(pending))) as pool:
                futures = {
                    pool.submit(run_case, name, asdict(config), str(output_dir)): name
                    for name, config in pending.items()
                }
                for future in as_completed(futures):
                    name = futures[future]
                    item = future.result()
                    print(
                        f"[{name}] crossing={item['drying_crossing_h']:.12f} h; "
                        f"R72={item['predicted_radius_final_cm']:.12f} cm; "
                        f"runtime={item['runtime_s']:.2f} s",
                        flush=True,
                    )
    missing = [
        name for name in CASES
        if not (output_dir / name / "result_full_precision.npz").exists()
    ]
    if missing:
        raise FileNotFoundError(f"cannot summarize; missing cases: {missing}")
    metadata = summarize(output_dir)
    print(json.dumps(metadata["attachment2_comparison"], ensure_ascii=False, indent=2))
    print(json.dumps(metadata["numerical_accuracy"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
